"""Saskatchewan's assisted flow in a real browser, against a stand-in registry served locally.

The stand-in behaves like the real one as far as it is known (its pstLookup.js): a terms page where a person
passes a CAPTCHA once per session, then a search page with a #pstSearch box. Here the "person" is a script on
the terms page that submits it after a moment. A second mode asks for the CAPTCHA again after every search, which
must stop bulk checking rather than keep asking.

Needs a browser (Playwright's Chromium, Edge, Chrome, or TAXCHECK_CHROMIUM); skipped otherwise.
"""

import datetime as dt
import html
import http.server
import tempfile
import threading
import unittest
import urllib.parse
from pathlib import Path
from unittest import mock

import openpyxl

from taxcheck import sk_pst
from taxcheck.browser import Browser, BrowserUnavailable
from taxcheck.checks import RowChecker
from taxcheck.result import ERROR, MANUAL, NOT_CONFIRMED, REGISTERED, Result

LICENSED = {"PRAIRIE HARDWARE LTD": "PSTVEND", "NORTHERN SUPPLY": "Registered Consumer",
            "SASKATCHEWAN TELECOMMUNICATIONS HOLDING CORPORATION": "PSTVEND"}

TERMS = """<html><body><h1>PST On-Line Registry</h1><form id="terms" method="post" action="/accept">
<input type="checkbox" id="termsChkBox"> I agree <div id="captcha_container">[I'm not a robot]</div>
<button id="termsNextBtn">Next</button></form>
<script>setTimeout(() => document.getElementById('terms').submit(), 400);</script></body></html>"""
# Like the real portal, every page has a site-wide search in its header; it must never be the one clicked.
HEADER = """<form method="get" action="/site-search"><input name="search_query"><button type="submit" title="Search">Search
</button></form>"""


def search_page(term="", result=""):
    """Laid out like the real page (2026-10-16): the search form echoes the term, its Search control is a
    javascript: link, and results follow the form. The submit is delayed, as a slow registry's would be, so a
    reader that doesn't wait for the new page reads the previous search's results."""
    return HEADER + f"""<html><body><h3>Instruction Box</h3>{result[0] if result else ""}
<p>The search ignores punctuation ... the number of searches in a 24-hour period is restricted to 100.</p>
<form id="pstSearchForm" method="post" action="/search">
<input id="pstSearch" name="pstSearch" value="{html.escape(term)}" maxlength="256">
<a href="javascript:submitPSTSearch('pstSearchForm')" class="btn btn-default search">Search</a></form>
<script>function submitPSTSearch(id) {{ setTimeout(() => document.getElementById(id).submit(), 300); }}</script>
{result[1] if result else ""}</body></html>"""


SEARCH = search_page()


def results_page(term):
    rows = [(n, kind) for n, kind in LICENSED.items() if term.upper().replace(" ", "") in n.replace(" ", "")]
    if not rows:
        return search_page(term, ('<div role="alert" class="alert alert-danger">No results found, please try a '
                                  'different search.</div>', ""))
    found = "Only 1 company found." if len(rows) == 1 else f"A total of {len(rows)} companies were found."
    table = ('<table id="pstLookUp" class="table dataTable table-bordered"><caption>Search Results</caption>'
             '<thead><tr><th scope="col">Business Name</th><th scope="col">Type</th></tr></thead><tbody>'
             + "".join(f"<tr><td>{html.escape(n)}</td><td>{k}</td></tr>" for n, k in rows) + "</tbody></table>")
    if FakeRegistry.datatables_empty:
        table += DATATABLES_EMPTY
    return search_page(term, ("", found + table))


# Seen on the user's PC (Edge, 2026-10-20): the page's DataTables reported no rows while the table showed three.
DATATABLES_EMPTY = """<script>window.jQuery = function (t) { return { DataTable: () => ({ rows: () => ({ data: () =>
({ toArray: () => [] }) }) }) }; };
window.jQuery.fn = { dataTable: { isDataTable: () => true } };</script>"""


ERROR_PAGE = """<html><body><h1>Technical Issue!</h1><p>We are having technical difficulties and are unable to
complete this action. Please try again later.</p></body></html>"""


class FakeRegistry(http.server.BaseHTTPRequestHandler):
    every_search = False  # True: the CAPTCHA is needed again after each search
    datatables_empty = False  # True: the page's table add-on reports no rows (the rows are in the page)
    expire_after = None  # the session ends after this many searches (once)
    broken = set()  # search terms the registry answers with an error page
    slow = {}  # search term -> seconds before it answers
    accepted = set()  # session ids that passed the terms page
    searches = []
    terms_shown = 0

    def log_message(self, *a):
        pass

    def _session(self):
        cookie = self.headers.get("Cookie", "")
        return cookie.split("sid=")[1].split(";")[0] if "sid=" in cookie else None

    def _send(self, body, cookie=None):
        self.send_response(200)
        self.send_header("Content-Type", "text/html")
        if cookie:
            self.send_header("Set-Cookie", f"sid={cookie}; Path=/")
        self.end_headers()
        self.wfile.write(body.encode())

    def do_GET(self):
        if self.path.startswith("/site-search"):
            FakeRegistry.searches.append("SITE SEARCH (wrong button)")
        if not self.path.startswith("/registry"):  # favicon etc.
            self.send_response(404)
            self.end_headers()
            return
        sid = self._session()
        if sid in FakeRegistry.accepted:
            self._send(SEARCH)
        else:
            FakeRegistry.terms_shown += 1
            self._send(TERMS, cookie=sid or f"s{FakeRegistry.terms_shown}")

    def do_POST(self):
        length = int(self.headers.get("Content-Length") or 0)
        form = urllib.parse.parse_qs(self.rfile.read(length).decode())
        sid = self._session()
        if self.path == "/accept":
            FakeRegistry.accepted.add(sid)
            self.send_response(303)
            self.send_header("Location", "/registry")
            self.end_headers()
            return
        if sid not in FakeRegistry.accepted:
            return self.do_GET()
        term = form.get("pstSearch", [""])[0]
        FakeRegistry.searches.append(term)
        if FakeRegistry.every_search:
            FakeRegistry.accepted.discard(sid)
        if FakeRegistry.expire_after is not None and len(FakeRegistry.searches) == FakeRegistry.expire_after:
            FakeRegistry.accepted.discard(sid)
            FakeRegistry.expire_after = None
            return self.do_GET()  # the terms page instead of results, as when a session times out
        if term.upper() in FakeRegistry.broken:
            self.send_response(500)
            self.send_header("Content-Type", "text/html")
            self.end_headers()
            self.wfile.write(ERROR_PAGE.encode())
            return
        if term.upper() in FakeRegistry.slow:
            import time
            time.sleep(FakeRegistry.slow[term.upper()])
        try:
            self._send(results_page(term))
        except OSError:  # the browser gave up on this answer
            pass


def browser_or_skip(test):
    b = Browser(debug_dir=test.tmp)
    try:
        b._ensure_started()
    except BrowserUnavailable as e:
        test.skipTest(f"no browser available: {e}")
    return b


class AssistedFlow(unittest.TestCase):
    def setUp(self):
        FakeRegistry.every_search, FakeRegistry.searches, FakeRegistry.terms_shown = False, [], 0
        FakeRegistry.datatables_empty = False
        FakeRegistry.accepted, FakeRegistry.broken, FakeRegistry.slow, FakeRegistry.expire_after = set(), set(), {}, None
        self.server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), FakeRegistry)
        threading.Thread(target=self.server.serve_forever, daemon=True).start()
        self.addCleanup(self.server.shutdown)
        url = f"http://127.0.0.1:{self.server.server_port}/registry"
        patch = mock.patch.object(sk_pst, "SK_REGISTRY", url)
        patch.start()
        self.addCleanup(patch.stop)
        # The browser must not send local requests through a proxy.
        env = mock.patch.dict("os.environ", {"HTTPS_PROXY": "", "HTTP_PROXY": ""})
        env.start()
        self.addCleanup(env.stop)
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.tmp = Path(self._tmp.name)
        self.browser = browser_or_skip(self)
        self.addCleanup(self.browser.close)

    def test_one_captcha_covers_every_search(self):
        sk = sk_pst.SkPstAssistant(self.browser, wait_minutes=0.5)
        self.assertEqual(sk.check(["PRAIRIE HARDWARE LTD"]).status, REGISTERED)
        res = sk.check(["NORTHERN SUPPLY"])
        self.assertEqual(res.status, REGISTERED)
        self.assertIn("registered consumer", res.detail)
        self.assertEqual(sk.check(["UNKNOWN WIDGETS"]).status, NOT_CONFIRMED)
        self.assertEqual(sk.captchas, 1)  # the person was asked once
        self.assertEqual(FakeRegistry.terms_shown, 1)
        self.assertEqual(FakeRegistry.searches, ["PRAIRIE HARDWARE", "NORTHERN SUPPLY", "UNKNOWN WIDGETS"])

    def test_each_search_reads_its_own_results(self):
        # Seen on a real run: the second search read the first one's results, because the results page loads
        # after the click returns. The stand-in delays its submit the same way.
        sk = sk_pst.SkPstAssistant(self.browser, wait_minutes=0.5)
        first = sk.check(["Prairie Hardware Ltd"])
        self.assertEqual(first.status, REGISTERED)
        self.assertEqual(first.registered_name, "PRAIRIE HARDWARE LTD | PSTVEND")
        sasktel = sk.check(["Saskatchewan Telecommunications"])
        self.assertEqual(sasktel.status, NOT_CONFIRMED)  # a similar name isn't this supplier's name
        self.assertIn("SASKATCHEWAN TELECOMMUNICATIONS HOLDING CORPORATION", sasktel.detail)
        self.assertNotIn("PRAIRIE", sasktel.detail)
        nothing = sk.check(["Zzqx Imaginary Widgets"])
        self.assertIn("found no business", nothing.detail)

    def test_saskatchewan_and_the_other_browser_checks_in_one_run(self):
        # Seen on a real run: the Saskatchewan window started first, then the hidden browser for Manitoba couldn't
        # start ("Playwright Sync API inside the asyncio loop"), so every Manitoba row became MANUAL CHECK.
        sk = sk_pst.SkPstAssistant(self.browser, wait_minutes=0.5)
        self.assertEqual(sk.check(["PRAIRIE HARDWARE LTD"]).status, REGISTERED)
        other = Browser(debug_dir=self.tmp)
        try:
            with other.page() as page:
                page.goto(sk_pst.SK_REGISTRY)
                self.assertTrue(page.content())
        finally:
            other.close()
        self.assertEqual(sk.check(["NORTHERN SUPPLY"]).status, REGISTERED)  # still works after the other closed

    def test_rows_read_from_the_page_when_the_table_add_on_reports_none(self):
        FakeRegistry.datatables_empty = True  # what the user's PC showed for FEDERATED CO-OPERATIVES
        sk = sk_pst.SkPstAssistant(self.browser, wait_minutes=0.5)
        res = sk.check(["PRAIRIE HARDWARE LTD"])
        self.assertEqual(res.status, REGISTERED, res.detail)
        self.assertIn("vendor's licence", res.detail)

    def test_a_resume_within_a_day_reuses_answers_without_the_browser(self):
        answers = self.tmp / "saskatchewan_answers.json"
        first = sk_pst.SkPstAssistant(self.browser, wait_minutes=0.5, answers_file=answers)
        self.assertEqual(first.check(["PRAIRIE HARDWARE LTD"]).status, REGISTERED)
        self.assertEqual(first.check(["UNKNOWN WIDGETS"]).status, NOT_CONFIRMED)
        searched = list(FakeRegistry.searches)

        class NoBrowser:  # the resumed run must not open the window for these
            debug_dir = None

            def new_persistent_page(self):
                raise AssertionError("no CAPTCHA needed")

        again = sk_pst.SkPstAssistant(NoBrowser(), answers_file=answers)
        self.assertTrue(again.answered(["PRAIRIE HARDWARE LTD"]))
        self.assertFalse(again.answered(["NORTHERN SUPPLY"]))
        self.assertTrue(again.answered(["AB"]))  # too short to search: nothing to ask
        res = again.check(["PRAIRIE HARDWARE LTD"])
        self.assertEqual(res.status, REGISTERED)
        self.assertIn("Registry answer from this tool's search on", res.detail)
        self.assertEqual(again.check(["UNKNOWN WIDGETS"]).status, NOT_CONFIRMED)
        self.assertEqual((FakeRegistry.searches, again.captchas), (searched, 0))

        with mock.patch.object(sk_pst.SkPstAssistant, "ANSWER_AGE", 0):  # a day later: searched again
            self.assertFalse(sk_pst.SkPstAssistant(NoBrowser(), answers_file=answers).answered(["PRAIRIE HARDWARE LTD"]))

    def test_probe_saves_each_page_and_what_was_read(self):
        import json

        import sk_probe
        from taxcheck import browser as browser_module

        real = browser_module.Browser
        with mock.patch.object(sk_probe, "HERE", self.tmp), \
                mock.patch.object(browser_module, "Browser", lambda headless=True, debug_dir=None: real(debug_dir=debug_dir)), \
                mock.patch("builtins.print"):
            self.assertEqual(sk_probe.main(["Prairie Hardware Ltd", "Zzqx Imaginary Widgets"]), 0)
        out = next((self.tmp / "logs").glob("sk_probe_*"))
        report = json.loads((out / "probe_report.json").read_text(encoding="utf-8"))
        self.assertEqual([r["tool_reads"] for r in report], [REGISTERED, NOT_CONFIRMED])
        self.assertEqual(len(list(out.glob("*.html"))), 2)
        self.assertEqual(len(list(out.glob("*.png"))), 2)
        import logging
        root = logging.getLogger()
        for h in list(root.handlers):
            root.removeHandler(h)
            h.close()

    def test_daily_search_limit(self):
        counter = self.tmp / "sk_searches.json"
        sk = sk_pst.SkPstAssistant(self.browser, wait_minutes=0.5, counter_file=counter, daily_limit=2)
        self.assertEqual(sk.check(["PRAIRIE HARDWARE LTD"]).status, REGISTERED)
        self.assertEqual(sk.check(["NORTHERN SUPPLY"]).status, REGISTERED)
        res = sk.check(["ANOTHER ONE LTD"])
        self.assertEqual(res.status, MANUAL)
        self.assertIn("allows 2 searches per 24 hours", res.detail)
        self.assertEqual(len(FakeRegistry.searches), 2)
        # The count is kept between runs.
        self.assertEqual(sk_pst.SkPstAssistant(self.browser, counter_file=counter, daily_limit=2).searches_left(), 0)

    def test_captcha_on_every_search_stops_bulk_checking(self):
        FakeRegistry.every_search = True
        sk = sk_pst.SkPstAssistant(self.browser, wait_minutes=0.5)
        self.assertEqual(sk.check(["PRAIRIE HARDWARE LTD"]).status, REGISTERED)
        # Asked again once (a session can just end), answered; asked again straight after: that's every search.
        self.assertEqual(sk.check(["NORTHERN SUPPLY"]).status, REGISTERED)
        res = sk.check(["SASKATCHEWAN TELECOMMUNICATIONS"])
        self.assertEqual(res.status, MANUAL)
        self.assertIn("asked for the CAPTCHA again", res.detail)
        self.assertEqual(sk.check(["ANOTHER ONE"]).status, MANUAL)  # no more asking
        self.assertEqual(sk.captchas, 2)
        self.assertEqual(len(FakeRegistry.searches), 2)

    def test_a_failed_search_doesnt_stop_the_others(self):
        FakeRegistry.broken = {"BROKEN WIDGETS"}
        sk = sk_pst.SkPstAssistant(self.browser, wait_minutes=0.5)
        self.assertEqual(sk.check(["PRAIRIE HARDWARE LTD"]).status, REGISTERED)
        res = sk.check(["BROKEN WIDGETS LTD"])
        self.assertEqual(res.status, ERROR)
        self.assertIn("failed twice", res.detail)
        self.assertEqual(FakeRegistry.searches.count("BROKEN WIDGETS"), 2)  # tried once more, then moved on
        self.assertIsNone(sk.unavailable)
        self.assertEqual(sk.check(["NORTHERN SUPPLY"]).status, REGISTERED)
        self.assertEqual(sk.check(["Saskatchewan Telecommunications"]).status, NOT_CONFIRMED)
        self.assertEqual(sk.captchas, 1)

    def test_a_search_with_no_answer_doesnt_stop_the_others_or_get_mixed_up(self):
        FakeRegistry.slow = {"SLOW SUPPLY": 4}
        sk = sk_pst.SkPstAssistant(self.browser, wait_minutes=0.5)
        with mock.patch.object(sk_pst, "LOAD_TIMEOUT_S", 1.5):
            self.assertEqual(sk.check(["SLOW SUPPLY"]).status, ERROR)
            res = sk.check(["PRAIRIE HARDWARE LTD"])
        self.assertEqual(res.status, REGISTERED)
        self.assertEqual(res.registered_name, "PRAIRIE HARDWARE LTD | PSTVEND")
        self.assertEqual(sk.check(["Zzqx Imaginary Widgets"]).status, NOT_CONFIRMED)
        self.assertIsNone(sk.unavailable)

    def test_session_ending_mid_run_asks_again_and_carries_on(self):
        FakeRegistry.expire_after = 2
        sk = sk_pst.SkPstAssistant(self.browser, wait_minutes=0.5)
        self.assertEqual(sk.check(["PRAIRIE HARDWARE LTD"]).status, REGISTERED)
        self.assertEqual(sk.check(["NORTHERN SUPPLY"]).status, REGISTERED)  # landed on the terms page, asked again
        self.assertEqual(sk.check(["Zzqx Imaginary Widgets"]).status, NOT_CONFIRMED)
        self.assertEqual(sk.check(["PRAIRIE HARDWARE LTD"]).status, REGISTERED)
        self.assertEqual(sk.captchas, 2)
        self.assertIsNone(sk.unavailable)

    def test_captcha_not_ticked_stops_without_waiting_again(self):
        sk = sk_pst.SkPstAssistant(self.browser, wait_minutes=0.5)
        sk._ready = mock.Mock(side_effect=sk_pst.CaptchaNotCompleted("the CAPTCHA wasn't completed in time"))
        res = sk.check(["PRAIRIE HARDWARE LTD"])
        self.assertEqual(res.status, MANUAL)
        self.assertIn("wasn't completed in time", res.detail)
        self.assertEqual(sk.check(["NORTHERN SUPPLY"]).status, MANUAL)
        self.assertEqual(sk._ready.call_count, 1)

    def test_window_closed_mid_run(self):
        sk = sk_pst.SkPstAssistant(self.browser, wait_minutes=0.5)
        self.assertEqual(sk.check(["PRAIRIE HARDWARE LTD"]).status, REGISTERED)
        sk.page.close()
        res = sk.check(["NORTHERN SUPPLY"])
        self.assertEqual(res.status, MANUAL)
        self.assertIn("window was closed", res.detail)

    def test_run_does_all_saskatchewan_searches_first_and_rows_reuse_them(self):
        import check_suppliers

        sk = sk_pst.SkPstAssistant(self.browser, wait_minutes=0.5)

        class Gst:
            def check(self, bn9, names, date, extra_names=()):
                return Result(REGISTERED, "fake CRA", names[0])

        checker = RowChecker(Gst(), None, sk=sk, available={"name", "gst", "province"})
        rows = [{"name": "PRAIRIE HARDWARE LTD", "gst": "108161779", "province": "SK"},
                {"name": "ACME LTD", "gst": "857305932", "province": "AB"},
                {"name": "NORTHERN SUPPLY", "gst": "108161779", "province": "SK"},
                {"name": "PRAIRIE HARDWARE LTD", "gst": "108161779", "province": "SK"}]  # duplicate: one search
        check_suppliers.prefetch_saskatchewan(checker, sk, rows)
        self.assertEqual(FakeRegistry.searches, ["PRAIRIE HARDWARE", "NORTHERN SUPPLY"])
        out = [checker.check(r, dt.date(2026, 9, 1)) for r in rows]
        self.assertEqual(FakeRegistry.searches, ["PRAIRIE HARDWARE", "NORTHERN SUPPLY"])  # nothing re-searched
        self.assertEqual(out[0]["sk_pst"].status, REGISTERED)
        self.assertNotIn("sk_pst", out[1])
        self.assertEqual(out[2]["sk_pst"].status, REGISTERED)

    def test_run_carries_on_past_a_failed_search(self):
        import check_suppliers

        FakeRegistry.broken = {"BROKEN WIDGETS"}
        sk = sk_pst.SkPstAssistant(self.browser, wait_minutes=0.5)
        checker = RowChecker(None, None, sk=sk, taxes={"sk_pst"}, available={"name", "province"})
        rows = [{"name": "PRAIRIE HARDWARE LTD", "province": "SK"},
                {"name": "BROKEN WIDGETS LTD", "province": "SK"},
                {"name": "NORTHERN SUPPLY", "province": "SK"},
                {"name": "ZZQX IMAGINARY WIDGETS", "province": "SK"}]
        check_suppliers.prefetch_saskatchewan(checker, sk, rows)
        self.assertEqual(FakeRegistry.searches, ["PRAIRIE HARDWARE", "BROKEN WIDGETS", "BROKEN WIDGETS",
                                                 "NORTHERN SUPPLY", "ZZQX IMAGINARY WIDGETS"])
        out = [checker.check(r, dt.date(2026, 9, 1))["sk_pst"].status for r in rows]
        self.assertEqual(out, [REGISTERED, ERROR, REGISTERED, NOT_CONFIRMED])


class UnreadableResults(unittest.TestCase):
    """The page says it found businesses but no rows were read: an ERROR to retry, never a cached answer."""

    PAGE = ('<form id="pstSearchForm"><input id="pstSearch" value="WIDGET"></form> A total of 2 companies were found.'
            '<table id="pstLookUp" class="table dataTable"><thead><tr><th>Business Name</th><th>Type</th></tr></thead>'
            '<tbody>\n<tr>\n<td>PRAIRIE   WIDGET LIMITED</td>\n<td>PSTVEND</td>\n</tr>'
            '<tr><td>PRAIRIE WIDGET &amp; CO</td><td>CONS</td></tr></tbody></table>')

    def test_rows_from_the_html(self):
        self.assertEqual(sk_pst.rows_from_html(self.PAGE), [["PRAIRIE WIDGET LIMITED", "PSTVEND"],
                                                            ["PRAIRIE WIDGET & CO", "CONS"]])
        self.assertEqual(sk_pst.found_count("A total of 2 companies were found."), 2)
        self.assertEqual(sk_pst.found_count("Only 1 company found."), 1)
        self.assertEqual(sk_pst.found_count("No results found, please try a different search."), 0)

    def test_found_but_unread_is_an_error_not_an_answer(self):
        status, detail, _ = sk_pst.classify([], "A total of 3 companies were found.", ["PRAIRIE WIDGET LTD"])
        self.assertEqual(status, ERROR)
        self.assertIn("couldn't be read", detail)
        status, _, _ = sk_pst.classify([], "No results found, please try a different search.", ["PRAIRIE WIDGET LTD"])
        self.assertEqual(status, NOT_CONFIRMED)

    def test_unreadable_answers_are_never_cached_or_reused(self):
        import json
        import tempfile
        import time

        with tempfile.TemporaryDirectory() as tmp:
            answers = Path(tmp) / "a.json"
            answers.write_text(json.dumps({"PRAIRIE WIDGET": {"at": time.time(), "rows": [],
                                                              "text": "A total of 3 companies were found."}}))
            sk = sk_pst.SkPstAssistant(None, answers_file=answers)  # an answer cached by the earlier version
            self.assertFalse(sk.answered(["PRAIRIE WIDGET LTD"]))
            sk._save_answer("OTHER WIDGET", [], "A total of 1 companies were found.")
            self.assertNotIn("OTHER WIDGET", sk._answers)


class RealResultsPage(unittest.TestCase):
    """What the real registry returned on 2026-10-16 (probe run), read the way the tool reads it."""

    FEDERATED = [["FEDERATED CO-OPERATIVES LTD", "PSTVEND"], ["FEDERATED CO-OPERATIVES       LIMITED", "PSTVEND"],
                 ["FEDERATED CO-OPERATIVES LIMITED", "PSTVEND"]]

    def test_vendor_licence_code(self):
        status, detail, row = sk_pst.classify(self.FEDERATED, "A total of 3 companies were found.",
                                              ["Federated Co-operatives Limited"])
        self.assertEqual(status, REGISTERED)
        self.assertIn("vendor's licence", detail)
        self.assertEqual(row, "FEDERATED CO-OPERATIVES LTD | PSTVEND")

    def test_similar_name_only(self):
        status, detail, _ = sk_pst.classify([["SASKATCHEWAN TELECOMMUNICATIONS HOLDING CORPORATION", "PSTVEND"]],
                                            "Only 1 company found.", ["Saskatchewan Telecommunications"])
        self.assertEqual(status, NOT_CONFIRMED)
        self.assertIn("similar names", detail)

    def test_no_results_and_too_many(self):
        self.assertIn("found no business", sk_pst.classify(
            [], "No results found, please try a different search.", ["Zzqx Imaginary Widgets"])[1])
        self.assertIn("More than 25", sk_pst.classify(
            [], "The search resulted in more than 25 companies.  Please refine your search criteria", ["Acme"])[1])

    def test_instructions_dont_look_like_a_limit_or_no_result(self):
        instructions = ("Results will display to a maximum of 25 matches and the number of searches in a 24-hour "
                        "period is restricted to 100.")
        self.assertEqual(sk_pst.classify([], instructions, ["Acme Widgets"])[1], "Couldn't read the registry's answer.")


class Menu(unittest.TestCase):
    def test_saskatchewan_question(self):
        import check_suppliers

        for answers, expected in ((["", ], True), (["y"], True), (["N"], False), (["maybe", "n"], False)):
            with mock.patch("builtins.input", side_effect=answers), mock.patch("builtins.print"):
                self.assertEqual(check_suppliers.ask_saskatchewan(), expected, answers)

    def test_only_asked_when_saskatchewan_is_chosen(self):
        import check_suppliers

        asked = []
        with mock.patch.object(check_suppliers, "ask_taxes", return_value={"gst"}), \
                mock.patch.object(check_suppliers, "ask_saskatchewan", side_effect=lambda: asked.append(1) or True), \
                mock.patch.object(check_suppliers, "ask_gst_dates", return_value="none"), \
                mock.patch.object(check_suppliers, "ask_paid_since", return_value=None), \
                mock.patch.object(check_suppliers, "run", return_value=None), \
                tempfile.TemporaryDirectory() as tmp:
            check_suppliers.main([str(Path(tmp) / "x.xlsx"), "--ask", "--log-dir", tmp])
            self.assertEqual(asked, [])  # GST/HST only: no Saskatchewan question
            with mock.patch.object(check_suppliers, "ask_taxes", return_value={"gst", "sk_pst"}):
                check_suppliers.main([str(Path(tmp) / "x.xlsx"), "--ask", "--log-dir", tmp])
            self.assertEqual(asked, [1])
        import logging
        root = logging.getLogger()
        for h in list(root.handlers):
            root.removeHandler(h)
            h.close()


class SaskatchewanSearchList(unittest.TestCase):
    def test_searches_needed_match_the_row_checks(self):
        checker = RowChecker(None, None, available={"name", "province", "pst", "country"})
        rows = [{"name": "(DNU) PRAIRIE HARDWARE LTD O/A PH", "province": "SK"},
                {"name": "ACME LTD", "province": "AB", "pst": "1234567"},
                {"name": "WESTERN LTD", "province": "SK", "country": "US"},
                {"name": "SASK CO", "province": "AB", "sk_pst": "7654321"}]
        jobs = checker.saskatchewan_searches(rows)
        self.assertEqual([j[0][:2] for j in jobs], [["PRAIRIE HARDWARE LTD O/A PH", "PRAIRIE HARDWARE LTD"],
                                                   ["SASK CO", "SASK COMPANY"]])
        self.assertEqual([j[1] for j in jobs], [False, True])
        self.assertEqual(RowChecker(None, None, taxes={"gst"}).saskatchewan_searches(rows), [])


if __name__ == "__main__":
    unittest.main()


class TheWholeSaskatchewanPlan(unittest.TestCase):
    """The user's plan end to end, with the real Saskatchewan code, a real browser and the stand-in registry:
    a full run with Saskatchewan (up to the daily limit), batches on the following days until all are answered,
    then Resume Last Run to put them in the full run's workbook, without a CAPTCHA or a search."""

    setUp = AssistedFlow.setUp

    def run_cli(self, *args, answers=()):
        import logging

        import check_suppliers

        root = logging.getLogger()
        for h in list(root.handlers):
            root.removeHandler(h)
            h.close()
        headless = lambda headless=True, debug_dir=None: Browser(debug_dir=debug_dir)  # no screen in tests
        with mock.patch.object(check_suppliers, "Browser", headless), \
                mock.patch.object(sk_pst, "DAILY_LIMIT", 2), \
                mock.patch.object(check_suppliers, "LAST_RUN", self.tmp / ".last_run"), \
                mock.patch("builtins.input", side_effect=list(answers)):
            code = check_suppliers.main([*args, "--log-dir", str(self.tmp / "logs"), "--cache-dir", str(self.tmp),
                                         "--no-name-lookup", "--no-browser", "--workers", "1"])
        for h in list(root.handlers):
            root.removeHandler(h)
            h.close()
        self.assertEqual(code, 0)

    def sk_column(self, name):
        ws = openpyxl.load_workbook(self.tmp / name).worksheets[0]
        headers = [c.value for c in ws[1]]
        return {r[0]: r[headers.index("SK PST Status")] for r in ws.iter_rows(min_row=2, values_only=True)
                if r[0] and r[headers.index("SK PST Status")]}

    def test_full_run_then_daily_batches_then_resume(self):
        import check_suppliers

        class Gst:  # GST/HST isn't what this test is about
            simplified = mock.Mock(error=None, _rows={})

            def __init__(self, *a, **k):
                pass

            def check(self, bn9, names, date, extra_names=()):
                return Result(REGISTERED, "fake CRA", names[0])

        src = self.tmp / "vendors.xlsx"
        wb = openpyxl.Workbook()
        for row in [["Company Name", "State", "Tax ID"], ["PRAIRIE HARDWARE LTD", "SK", "108161779RT0001"],
                    ["NORTHERN SUPPLY", "SK", "108161779RT0001"], ["ACME LTD", "ON", "857305932RT0001"],
                    ["ZZQX IMAGINARY WIDGETS", "SK", "108161779RT0001"]]:
            wb.active.append(row)
        wb.save(src)
        with mock.patch.object(check_suppliers, "GstChecker", Gst), \
                mock.patch.object(check_suppliers, "SkPstAssistant",
                                  lambda b, **kw: sk_pst.SkPstAssistant(b, wait_minutes=0.5, daily_limit=2,
                                                                        **{k: v for k, v in kw.items() if k != "daily_limit"})):
            # Night 1: the full run, Saskatchewan Y; the daily limit (2 here) stops the third search.
            self.run_cli(str(src), "--sk-assist", "--gst-dates", "none")
            self.assertEqual(FakeRegistry.searches, ["PRAIRIE HARDWARE", "NORTHERN SUPPLY"])
            self.assertEqual(self.sk_column("vendors_tax_check.xlsx")["ZZQX IMAGINARY WIDGETS"], "MANUAL CHECK")
            shown = FakeRegistry.terms_shown

            # Same day, a batch: nothing left to search today, so no window, no CAPTCHA.
            self.run_cli("--sk-only")
            self.assertEqual(FakeRegistry.terms_shown, shown)
            self.assertEqual(self.sk_column("vendors_saskatchewan.xlsx")["PRAIRIE HARDWARE LTD"], "REGISTERED")

            # Next day (the 24 hours have passed): the batch searches only what's left.
            (self.tmp / "saskatchewan_searches.json").write_text("[]")
            self.run_cli("--sk-only")
            self.assertEqual(FakeRegistry.searches, ["PRAIRIE HARDWARE", "NORTHERN SUPPLY", "ZZQX IMAGINARY WIDGETS"])
            batch = self.sk_column("vendors_saskatchewan.xlsx")
            self.assertEqual(batch["ZZQX IMAGINARY WIDGETS"], "NOT CONFIRMED")
            self.assertNotIn("ACME LTD", batch)

            # All answered: Resume Last Run puts them in the full run's workbook, with no CAPTCHA and no search.
            shown, searched = FakeRegistry.terms_shown, list(FakeRegistry.searches)
            self.run_cli("--resume", "last")
            self.assertEqual((FakeRegistry.terms_shown, FakeRegistry.searches), (shown, searched))
            final = self.sk_column("vendors_tax_check.xlsx")
            self.assertEqual(final, {"PRAIRIE HARDWARE LTD": "REGISTERED", "NORTHERN SUPPLY": "REGISTERED",
                                     "ZZQX IMAGINARY WIDGETS": "NOT CONFIRMED"})
