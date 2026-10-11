"""Reading messy workbooks and writing results back."""

import tempfile
import unittest
from pathlib import Path

import openpyxl

from taxcheck.result import MANUAL, NOT_CONFIRMED, REGISTERED, Result
from taxcheck.workbook import SupplierSheet, clean_value


def make_workbook(path, rows):
    wb = openpyxl.Workbook()
    for row in rows:
        wb.active.append(row)
    wb.save(path)


class WorkbookIO(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_numbers_typed_as_numbers(self):
        self.assertEqual(clean_value(857305932.0), 857305932)
        self.assertEqual(clean_value("  857305932RT0001 "), "857305932RT0001")
        self.assertIsNone(clean_value("   "))
        path = self.dir / "in.xlsx"
        make_workbook(path, [["Supplier Name", "GST/HST Number", "QST Number"], ["Amazon.com.ca ULC", 857305932.0, 1019288451]])
        (_, values), = SupplierSheet(path).rows()
        self.assertEqual(values["gst"], 857305932)
        self.assertEqual(values["qst"], 1019288451)

    def test_output_keeps_original_columns_and_formulas(self):
        path, out = self.dir / "in.xlsx", self.dir / "out.xlsx"
        make_workbook(path, [["Supplier Name", "GST/HST Number", "Note"], ["A Co", "108161779RT0001", "=1+1"]])
        sheet = SupplierSheet(path)
        sheet.write_results({2: {"gst": Result(REGISTERED, "ok", "A CO")}}, out, ["note"])
        ws = openpyxl.load_workbook(out)["Sheet"]
        headers = [c.value for c in ws[1]]
        self.assertEqual(headers[:4], ["Supplier Name", "GST/HST Number", "Note", "Overall"])
        self.assertEqual(ws["C2"].value, "=1+1")
        row = dict(zip(headers, [c.value for c in ws[2]]))
        self.assertEqual(row["Overall"], "OK")
        self.assertEqual(row["GST/HST Status"], REGISTERED)
        self.assertEqual(row["Name on Government Record"], "A CO")
        self.assertIn("Tax Check Summary", openpyxl.load_workbook(out).sheetnames)

    def test_rerun_on_results_file_overwrites_result_columns(self):
        path, out1, out2 = self.dir / "in.xlsx", self.dir / "r1.xlsx", self.dir / "r2.xlsx"
        make_workbook(path, [["Supplier Name", "GST/HST Number", "QST Number"], ["A Co", "108161779", "1019288451TQ0004"]])
        SupplierSheet(path).write_results(
            {2: {"gst": Result(NOT_CONFIRMED, "first"), "qst": Result(REGISTERED, "q")}}, out1, []
        )
        SupplierSheet(out1).write_results({2: {"gst": Result(REGISTERED, "second")}}, out2, [])
        wb = openpyxl.load_workbook(out2)
        ws = wb["Sheet"]
        headers = [c.value for c in ws[1]]
        self.assertEqual(headers.count("Overall"), 1)
        row = dict(zip(headers, [c.value for c in ws[2]]))
        self.assertEqual(row["GST/HST Details"], "second")
        self.assertIsNone(row["QST Status"])  # stale value from the first run is cleared
        self.assertEqual(wb.sheetnames.count("Tax Check Summary"), 1)

    def test_external_text_never_becomes_a_formula(self):
        path, out = self.dir / "in.xlsx", self.dir / "out.xlsx"
        make_workbook(path, [["Supplier Name", "GST/HST Number"], ["A Co", "108161779"]])
        SupplierSheet(path).write_results({2: {"gst": Result(REGISTERED, "=HYPERLINK(\"x\")", "=1+1")}}, out, [])
        ws = openpyxl.load_workbook(out)["Sheet"]
        headers = [c.value for c in ws[1]]
        row = dict(zip(headers, [c.value for c in ws[2]]))
        self.assertEqual(row["Name on Government Record"], "'=1+1")
        self.assertTrue(row["GST/HST Details"].startswith("'="))

    def test_overall_review_when_anything_not_registered(self):
        path, out = self.dir / "in.xlsx", self.dir / "out.xlsx"
        make_workbook(path, [["Supplier Name", "Province"], ["A Co", "SK"]])
        SupplierSheet(path).write_results(
            {2: {"gst": Result(REGISTERED, ""), "sk_pst": Result(MANUAL, "")}}, out, []
        )
        ws = openpyxl.load_workbook(out)["Sheet"]
        self.assertEqual(ws["C2"].value, "REVIEW")

    def test_no_header_row_is_a_clear_error(self):
        path = self.dir / "in.xlsx"
        make_workbook(path, [["foo", "bar"], [1, 2]])
        with self.assertRaisesRegex(ValueError, "header row"):
            SupplierSheet(path)

    def test_named_sheet(self):
        path = self.dir / "in.xlsx"
        wb = openpyxl.Workbook()
        wb.active.title = "Cover"
        ws = wb.create_sheet("Vendors")
        ws.append(["Vendor", "HST #"])
        ws.append(["A Co", "108161779"])
        wb.save(path)
        self.assertEqual(len(list(SupplierSheet(path, "Vendors").rows())), 1)


class InputFiles(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_csv_utf8_and_windows_encodings(self):
        for encoding in ("utf-8-sig", "cp1252"):
            path = self.dir / f"suppliers_{encoding}.csv"
            path.write_text("Supplier Name,Province,GST/HST Number\nCafé Québécois Inc,QC,857305932\n\n", encoding=encoding)
            sheet = SupplierSheet(path)
            (_, values), = sheet.rows()
            self.assertEqual(values["name"], "Café Québécois Inc")
            self.assertEqual(values["gst"], "857305932")
            out = self.dir / f"out_{encoding}.xlsx"
            sheet.write_results({2: {"gst": Result(REGISTERED, "")}}, out, [])
            self.assertEqual(openpyxl.load_workbook(out).active["D2"].value, "OK")

    def test_unsupported_files_get_plain_messages(self):
        from taxcheck.workbook import InputError

        (self.dir / "old.xls").write_bytes(b"x")
        with self.assertRaisesRegex(InputError, "Save As"):
            SupplierSheet(self.dir / "old.xls")
        with self.assertRaisesRegex(InputError, "Can't find"):
            SupplierSheet(self.dir / "nope.xlsx")
        make_workbook(self.dir / "in.xlsx", [["Supplier Name", "GST/HST Number"]])
        with self.assertRaisesRegex(InputError, "Sheets in this file: Sheet"):
            SupplierSheet(self.dir / "in.xlsx", "Vendors")


class DatesAndHeaders(unittest.TestCase):
    def test_common_date_formats(self):
        import datetime as dt

        from taxcheck.workbook import parse_date

        want = dt.date(2026, 9, 15)
        for text in ("2026-09-15", "2026/09/15", "15/09/2026", "15-Sep-2026", "15 Sep 2026", "September 15, 2026",
                     "Sep 15 2026", "2026-09-15 00:00:00", "2026-09-15T13:45:00", "20260915", "15-09-2026"):
            self.assertEqual(parse_date(text), want, text)
        self.assertEqual(parse_date("09/10/2026"), dt.date(2026, 10, 9))  # ambiguous: day first
        self.assertIsNone(parse_date("next tuesday"))

    def test_suggestions_for_near_miss_headers(self):
        from taxcheck.workbook import suggest_field

        self.assertEqual(suggest_field("QC Tax No")[1], "qst")
        self.assertEqual(suggest_field("GST Reg No.")[1], "gst")
        self.assertIsNone(suggest_field("Amount"))

    def test_common_company_headers_are_recognized(self):
        from taxcheck.workbook import _ALIASES, _key

        for header, field in [("Legal Entity", "name"), ("DBA Name", "trade_name"), ("State/Prov", "province"),
                              ("Operating As", "trade_name"), ("GST #", "gst"), ("QST No", "qst")]:
            self.assertEqual(_ALIASES.get(_key(header)), field, header)


class ScientificNotation(unittest.TestCase):
    def test_excel_mangled_numbers_are_flagged(self):
        from taxcheck.numbers import BadNumber, parse_gst, parse_qst

        for bad in ("8.57306E+08", "1.01929E+09", "8E+08"):
            with self.assertRaisesRegex(BadNumber, "scientific notation"):
                (parse_qst if bad.startswith("1.0") else parse_gst)(bad)


if __name__ == "__main__":
    unittest.main()
