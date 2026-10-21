"""The command line end to end (with fake government services), including the run log."""

import logging
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import openpyxl

import check_suppliers
from taxcheck.browser import BrowserUnavailable
from taxcheck.result import REGISTERED, Result


class FakeGst:
    def __init__(self, *a, **k):
        self.simplified = mock.Mock(error=None)

    def check(self, bn9, names, date):
        if bn9 == "857305932":
            raise ConnectionError("simulated network failure")
        return Result(REGISTERED, "fake CRA", names[0])


class FakeQst:
    nr_error = None

    def __init__(self, *a, **k):
        pass

    def check(self, number, suffix_assumed=False, judge=None):
        return Result(REGISTERED, "fake RQ")


class Cli(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.src = self.dir / "suppliers.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        for row in [
            ["Supplier Name", "Province", "GST/HST Number", "QST Number", "PST Number", "Transaction Date"],
            ["Good Co", "QC", "108161779RT0001", "1019288451TQ0004", None, "2026-09-30"],
            ["Flaky Co", "ON", "857305932RT0001", None, None, "not a date"],
            ["BC Co", "BC", "108161779RT0001", None, "PST-1000-7572", None],
            ["Typo Co", "ON", "123456789", None, None, None],
        ]:
            ws.append(row)
        wb.save(self.src)
        patches = [mock.patch.object(check_suppliers, "GstChecker", FakeGst),
                   mock.patch.object(check_suppliers, "QstChecker", FakeQst)]
        for p in patches:
            p.start()
            self.addCleanup(p.stop)

    def tearDown(self):
        # Detach the run's log handlers so the log file is closed before the temp folder is removed.
        root = logging.getLogger()
        for h in list(root.handlers):
            root.removeHandler(h)
            h.close()
        self.tmp.cleanup()

    def run_cli(self, *extra):
        logs = self.dir / "logs"
        code = check_suppliers.main([str(self.src), "--log-dir", str(logs), "--no-browser", "--no-name-lookup",
                                    "--workers", "1", *extra])
        log_files = list(logs.glob("tax_check_*.log"))
        self.assertEqual(len(log_files), 1)
        return code, log_files[0].read_text(encoding="utf-8")

    def test_full_run_writes_results_and_log(self):
        code, log_text = self.run_cli()
        self.assertEqual(code, 0)
        out = self.dir / "suppliers_tax_check.xlsx"
        ws = openpyxl.load_workbook(out)["Sheet"]
        headers = [c.value for c in ws[1]]
        rows = [dict(zip(headers, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
        self.assertEqual([r["Overall"] for r in rows], ["OK", "REVIEW", "REVIEW", "REVIEW"])
        self.assertEqual(rows[1]["GST/HST Status"], "ERROR")         # network failure is contained
        self.assertEqual(rows[2]["BC PST Status"], "MANUAL CHECK")   # --no-browser
        self.assertEqual(rows[3]["GST/HST Status"], "INVALID NUMBER")

        # The log has what's needed to debug afterwards.
        self.assertIn("Arguments:", log_text)
        self.assertIn("column gst", log_text)
        self.assertIn("---- Row 3: Flaky Co", log_text)
        self.assertIn("Traceback", log_text)
        self.assertIn("simulated network failure", log_text)
        self.assertIn("couldn't read date 'not a date'", log_text)
        self.assertIn("Results written to", log_text)

    def test_ctrl_c_keeps_results_so_far(self):
        calls = []

        def check(self_, bn9, names, date):
            calls.append(bn9)
            if len(calls) == 2:
                raise KeyboardInterrupt
            return Result(REGISTERED, "fake CRA", names[0])

        with mock.patch.object(FakeGst, "check", check):
            code, log_text = self.run_cli()
        self.assertEqual(code, 0)
        wb = openpyxl.load_workbook(self.dir / "suppliers_tax_check.xlsx")
        self.assertEqual(wb["Sheet"]["G2"].value, "OK")  # first supplier's result was kept
        notes = [r[0] for r in wb["Tax Check Summary"].iter_rows(values_only=True) if r and r[0]]
        self.assertTrue(any(str(n).startswith("STOPPED EARLY") for n in notes), notes)

    def test_stop_after_then_resume_last(self):
        # What Run Test.bat does: stop as Ctrl+C would, then resume through the "last run" note.
        with mock.patch.object(check_suppliers, "LAST_RUN", self.dir / ".last_run"):
            code, log_text = self.run_cli("--stop-after", "2")
            self.assertEqual(code, 0)
            self.assertIn("Stopped early after 2 of 4", log_text)
            for h in list(logging.getLogger().handlers):  # a second run in the same test: start a fresh log
                logging.getLogger().removeHandler(h)
                h.close()
            with mock.patch("builtins.input", return_value="n") as asked:  # Saskatchewan was skipped: asked again
                code = check_suppliers.main(["--resume", "last", "--log-dir", str(self.dir / "logs2"), "--workers", "1"])
            asked.assert_called_once()
        self.assertEqual(code, 0)
        ws = openpyxl.load_workbook(self.dir / "suppliers_tax_check.xlsx")["Sheet"]
        headers = [c.value for c in ws[1]]
        self.assertTrue(all(r[headers.index("Overall")] for r in ws.iter_rows(min_row=2, values_only=True)))
        resumed = (self.dir / "logs2").glob("tax_check_*.log").__next__().read_text(encoding="utf-8")
        # Of the 2 before the stop, the one whose lookup failed (Flaky Co) is checked again.
        self.assertIn("1 suppliers already done, 1 with errors", resumed)
        self.assertIn("'no_browser': True", resumed)  # the stopped run's options, reused

    def test_check_columns_is_offline_and_explains_mapping(self):
        with mock.patch.object(FakeGst, "check", side_effect=AssertionError("no lookups in --check-columns")):
            code, log_text = self.run_cli("--check-columns", "--preview", "2")
        self.assertEqual(code, 0)
        self.assertIn("'GST/HST Number' -> gst", log_text)
        self.assertIn("WOULD CHECK", log_text)
        self.assertIn("What the first 2 would check", log_text)
        self.assertFalse((self.dir / "suppliers_tax_check.xlsx").exists())

    def test_check_columns_on_unrecognized_headers_lists_them(self):
        wb = openpyxl.Workbook()
        wb.active.append(["Vendor ID", "Entity", "Tax Reg #"])
        wb.save(self.src)
        code, log_text = self.run_cli("--check-columns")
        self.assertEqual(code, 1)
        self.assertIn("'Tax Reg #'", log_text)
        self.assertIn("COLUMNS", log_text)

    def test_no_path_opens_picker_and_outputs_go_next_to_workbook(self):
        with mock.patch.object(check_suppliers, "choose_workbook", return_value=str(self.src)) as picker:
            code = check_suppliers.main(["--no-browser", "--no-name-lookup"])
        self.assertEqual(code, 0)
        picker.assert_called_once()
        self.assertTrue((self.dir / "suppliers_tax_check.xlsx").exists())
        self.assertEqual(len(list((self.dir / "logs").glob("tax_check_*.log"))), 1)  # logs beside the workbook

    def test_cancelled_picker_does_nothing(self):
        with mock.patch.object(check_suppliers, "choose_workbook", return_value=None):
            self.assertEqual(check_suppliers.main(["--no-browser"]), 1)
        self.assertFalse((self.dir / "suppliers_tax_check.xlsx").exists())

    def test_picker_falls_back_to_typed_path_without_a_screen(self):
        import builtins

        with mock.patch.dict("sys.modules", {"tkinter": None}), \
                mock.patch.object(builtins, "input", return_value=f'"{self.src}"'):
            self.assertEqual(check_suppliers.choose_workbook(), str(self.src))  # quotes from "Copy as path" removed

    def test_missing_file_fails_cleanly_with_log(self):
        self.src = self.dir / "nope.xlsx"
        code, log_text = self.run_cli()
        self.assertEqual(code, 1)
        self.assertIn("Can't find the input file", log_text)
        self.assertNotIn("Traceback", log_text)  # a missing file is the user's to fix, not a crash

    def test_browser_start_failure_falls_back_to_manual(self):
        with mock.patch(
            "taxcheck.browser.Browser._ensure_started", side_effect=BrowserUnavailable("no chromium. Run playwright install")
        ):
            code = check_suppliers.main([str(self.src), "--log-dir", str(self.dir / "logs2"), "--no-name-lookup"])
        self.assertEqual(code, 0)
        ws = openpyxl.load_workbook(self.dir / "suppliers_tax_check.xlsx")["Sheet"]
        headers = [c.value for c in ws[1]]
        row = dict(zip(headers, [c.value for c in ws[4]]))
        self.assertEqual(row["BC PST Status"], "MANUAL CHECK")
        self.assertIn("no chromium", row["BC PST Details"])

if __name__ == "__main__":
    unittest.main()


class ResumeSaskatchewan(unittest.TestCase):
    """--resume checks again Saskatchewan rows whose searches couldn't run (the daily limit...), as the run says."""

    def test_rows_left_for_a_manual_check_by_the_searches_are_redone(self):
        import csv
        import json
        import tempfile
        from pathlib import Path

        from taxcheck.progress import load_resume

        rows = {
            2: {"gst": ["REGISTERED", "ok", ""], "sk_pst": ["REGISTERED", "vendor's licence", ""]},
            3: {"gst": ["REGISTERED", "ok", ""], "sk_pst": ["MANUAL CHECK", "Saskatchewan's registry allows 100 "
                                                             "searches per 24 hours and they're used up.", ""]},
            4: {"gst": ["REGISTERED", "ok", ""], "sk_pst": ["MANUAL CHECK", "needs at least 4 letters of the "
                                                             "business name", ""]},
            5: {"gst": ["ERROR", "down", ""]},
        }
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "progress.csv"
            with open(path, "w", newline="", encoding="utf-8-sig") as f:
                w = csv.DictWriter(f, fieldnames=["sheet_row", "results_json"])
                w.writeheader()
                for r, data in rows.items():
                    w.writerow({"sheet_row": r, "results_json": json.dumps(data)})
            self.assertEqual(sorted(load_resume(path)), [2, 4])  # 3 (limit) and 5 (error) are checked again


class ResumeLastRun(unittest.TestCase):
    """"Resume Last Run.bat" (--resume last): the latest full run's workbook, progress file and answers."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        patcher = mock.patch.object(check_suppliers, "LAST_RUN", Path(self.tmp.name) / ".last_run")
        patcher.start()
        self.addCleanup(patcher.stop)
        self.progress = Path(self.tmp.name) / "progress_20260101_000000.csv"
        self.progress.write_text("sheet_row,results_json\n", encoding="utf-8")

    def args(self, **kw):
        base = dict(limit=None, offline=False, workbook=str(Path(self.tmp.name) / "vendors.xlsx"), taxes={"gst", "qst"},
                    sk_assist=False, gst_dates="none", paid_since="2022-01-01")
        base.update(kw)
        return mock.Mock(**{k: base.get(k) for k in check_suppliers.RUN_SETTINGS + ("limit", "offline", "workbook")})

    def resumed_args(self, answers=()):
        seen = {}
        with mock.patch.object(check_suppliers, "run", side_effect=lambda args, log_path: seen.update(vars(args))), \
                mock.patch.object(check_suppliers, "setup_logging", return_value=Path(self.tmp.name) / "x.log"), \
                mock.patch.object(check_suppliers, "apply_tailoring", return_value=""), \
                mock.patch("builtins.input", side_effect=list(answers)) as asked:
            code = check_suppliers.main(["--resume", "last"])
        return code, seen, asked

    def test_same_workbook_progress_and_answers_without_asking(self):
        check_suppliers.remember_run(self.args(), self.progress)
        code, seen, asked = self.resumed_args()
        self.assertEqual(code, 0)
        self.assertEqual(seen["resume"], str(self.progress.resolve()))
        self.assertEqual(Path(seen["workbook"]).name, "vendors.xlsx")
        self.assertEqual(seen["taxes"], {"gst", "qst"})
        self.assertEqual((seen["gst_dates"], seen["paid_since"], seen["ask"]), ("none", "2022-01-01", False))
        asked.assert_not_called()  # Saskatchewan wasn't among the taxes: nothing to ask

    def test_saskatchewan_skipped_last_time_is_asked_again(self):
        check_suppliers.remember_run(self.args(taxes=None), self.progress)
        code, seen, asked = self.resumed_args(answers=["y"])
        self.assertEqual(code, 0)
        self.assertTrue(seen["sk_assist"])
        self.assertIsNone(seen["taxes"])  # everything, as before

    def test_trials_and_offline_runs_arent_resumed(self):
        check_suppliers.remember_run(self.args(limit=25), self.progress)
        check_suppliers.remember_run(self.args(offline=True), self.progress)
        code, seen, _ = self.resumed_args()
        self.assertEqual((code, seen), (1, {}))  # nothing to resume: says so instead of starting a run

    def test_missing_progress_file(self):
        check_suppliers.remember_run(self.args(), self.progress)
        self.progress.unlink()
        self.assertEqual(self.resumed_args()[0], 1)


class TestRunIsolation(unittest.TestCase):
    """Run Test.bat (TAXCHECK_TEST_RUN) must never resume or overwrite the user's own run, or use their tailoring."""

    def test_test_runs_keep_their_own_last_run_note_and_skip_tailoring(self):
        import os
        import subprocess
        import sys

        with tempfile.TemporaryDirectory() as tmp:
            code = ("import check_suppliers as c; print(c.LAST_RUN); print(bool(c.TEST_RUN)); "
                    "c.remember_workbook('x.xlsx'); print((c.HERE / '.last_workbook').exists())")
            env = dict(os.environ, TAXCHECK_TEST_RUN=tmp)
            before = (check_suppliers.HERE / ".last_workbook").exists()
            out = subprocess.run([sys.executable, "-c", code], cwd=check_suppliers.HERE, env=env, text=True,
                                 capture_output=True, check=True).stdout.split("\n")
        self.assertEqual(Path(out[0]), Path(tmp) / ".last_run")
        self.assertEqual(out[1], "True")
        self.assertEqual(out[2], str(before))  # .last_workbook left as it was


class SaskatchewanBatches(unittest.TestCase):
    """"Run Saskatchewan Batch.bat" (--sk-only): only Saskatchewan suppliers, up to the daily limit, then the rest
    on later days; answers are kept, so each batch searches only what's left."""

    LIMIT = 2

    def setUp(self):
        import json

        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.dir = Path(self.tmp.name)
        self.src = self.dir / "vendors.xlsx"
        wb = openpyxl.Workbook()
        for row in [["Company Name", "State", "Tax ID"], ["PRAIRIE WIDGETS LTD", "SK", "108161779RT0001"],
                    ["ACME LTD", "ON", "857305932RT0001"], ["NORTHERN SUPPLY CO", "SK", ""],
                    ["REGINA PAINT INC", "Sask.", ""], ["GOOD CO", "BC", ""]]:
            wb.active.append(row)
        wb.save(self.src)
        self.searched = []
        test = self

        class FakeSk:  # the registry: LIMIT searches per run ("per day"), answers saved between runs
            daily_limit = test.LIMIT

            def __init__(self, browser, counter_file=None, answers_file=None, **kw):
                self.file, self.left, self.captchas, self.unavailable = answers_file, test.LIMIT, 0, None
                self.saved = json.loads(answers_file.read_text()) if answers_file.exists() else {}

            def answered(self, names):
                return names[0] in self.saved

            def searches_left(self):
                return self.left

            def _limit_message(self):
                return ("Saskatchewan's registry allows 2 searches per 24 hours and they're used up. The next full "
                        "batch is possible from 2026-10-22 09:00 (a few searches from 08:55). Double-click later.")

            def check(self, names):
                if names[0] in self.saved:
                    return Result(REGISTERED, "vendor's licence (Registry answer from this tool's search on ...)")
                if self.unavailable or not self.left:
                    self.unavailable = self._limit_message()
                    return Result("MANUAL CHECK", self.unavailable)
                self.left -= 1
                test.searched.append(names[0])
                self.saved[names[0]] = True
                self.file.write_text(json.dumps(self.saved))
                return Result(REGISTERED, "vendor's licence")

        class NoGst(FakeGst):
            def check(self, *a, **k):
                raise AssertionError("a Saskatchewan batch never looks up GST/HST")

        for p in (mock.patch.object(check_suppliers, "SkPstAssistant", FakeSk),
                  mock.patch.object(check_suppliers, "GstChecker", NoGst),
                  mock.patch.object(check_suppliers, "QstChecker", FakeQst),
                  mock.patch.object(check_suppliers, "LAST_RUN", self.dir / ".last_run")):
            p.start()
            self.addCleanup(p.stop)
        (self.dir / ".last_run").write_text('{"workbook": "%s", "progress": "x"}' % str(self.src).replace("\\", "/"))

    def tearDown(self):
        root = logging.getLogger()
        for h in list(root.handlers):
            root.removeHandler(h)
            h.close()

    def batch(self):
        for h in list(logging.getLogger().handlers):
            logging.getLogger().removeHandler(h)
            h.close()
        code = check_suppliers.main(["--sk-only", "--log-dir", str(self.dir / "logs"), "--cache-dir", str(self.dir)])
        self.assertEqual(code, 0)
        result = openpyxl.load_workbook(self.dir / "vendors_saskatchewan.xlsx")
        ws = result.worksheets[0]
        headers = [c.value for c in ws[1]]
        rows = {r[0]: dict(zip(headers, r)) for r in ws.iter_rows(min_row=2, values_only=True)}
        notes = [r[0] for r in result["Tax Check Summary"].iter_rows(values_only=True) if r and r[0]]
        return rows, next(n for n in notes if str(n).startswith("SASKATCHEWAN BATCH"))

    def test_day_by_day_until_done(self):
        before = (self.dir / ".last_run").read_text()
        rows, note = self.batch()  # day 1: the workbook comes from the last full run, no questions
        self.assertEqual(self.searched, ["PRAIRIE WIDGETS LTD", "NORTHERN SUPPLY CO"])
        self.assertIn("2 of 3 Saskatchewan suppliers have an answer; 1 still to search", note)
        self.assertIn("The next full batch is possible from 2026-10-22 09:00", note)
        self.assertEqual(rows["REGINA PAINT INC"]["SK PST Status"], "MANUAL CHECK")
        self.assertFalse(rows["ACME LTD"].get("SK PST Status"))  # not a Saskatchewan supplier: untouched
        self.assertNotIn("GST/HST Status", rows["PRAIRIE WIDGETS LTD"])  # only Saskatchewan columns

        rows, note = self.batch()  # day 2: only what's left is searched
        self.assertEqual(self.searched, ["PRAIRIE WIDGETS LTD", "NORTHERN SUPPLY CO", "REGINA PAINT INC"])
        self.assertIn("all 3 Saskatchewan suppliers are done: 3 have an answer", note)
        self.assertIn("Resume Last Run.bat", note)
        self.assertEqual({r["SK PST Status"] for n, r in rows.items() if n != "ACME LTD" and n != "GOOD CO"},
                         {"REGISTERED"})
        self.assertEqual((self.dir / ".last_run").read_text(), before)  # the full run's note is never replaced

    def test_the_full_runs_paid_since_cut_off_carries_over(self):
        import json

        (self.dir / ".last_run").write_text(json.dumps({"workbook": str(self.src), "progress": "x",
                                                        "paid_since": "2024-01-01"}))
        seen = {}
        with mock.patch.object(check_suppliers, "run", side_effect=lambda args, log_path: seen.update(vars(args))):
            check_suppliers.main(["--sk-only", "--log-dir", str(self.dir / "logs"), "--cache-dir", str(self.dir)])
        self.assertEqual((seen["paid_since"], seen["taxes"], seen["sk_assist"]), ("2024-01-01", {"sk_pst"}, True))


class SaskatchewanBatchSummary(unittest.TestCase):
    def test_never_all_done_after_a_stop(self):
        sk = mock.Mock(unavailable=None)
        line = check_suppliers.saskatchewan_batch_summary({}, sk, 946, stopped_early=True)
        self.assertIn("0 of 946 Saskatchewan suppliers have an answer; 946 still to search", line)
        self.assertIn("stopped before the end", line)
        self.assertNotIn("Resume Last Run", line)

    def test_counts_and_by_hand(self):
        from taxcheck.result import ERROR, MANUAL, NOT_CONFIRMED

        results = {2: {"sk_pst": Result(REGISTERED, "licence")}, 3: {"sk_pst": Result(NOT_CONFIRMED, "no business")},
                   4: {"sk_pst": Result(MANUAL, "needs at least 4 letters of the business name")},
                   5: {"sk_pst": Result(ERROR, "failed twice")},
                   6: {"sk_pst": Result(MANUAL, "Saskatchewan's registry didn't answer this search in 2 batches")}}
        line = check_suppliers.saskatchewan_batch_summary(results, mock.Mock(unavailable=None), 6)
        self.assertIn("2 of 6 Saskatchewan suppliers have an answer; 2 still to search. 2 need a search by hand", line)
        done = {2: results[2], 3: results[3], 4: results[4]}
        self.assertIn("all 3 Saskatchewan suppliers are done: 2 have an answer. 1 need a search by hand",
                      check_suppliers.saskatchewan_batch_summary(done, None, 3))
