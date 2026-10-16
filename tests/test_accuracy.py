"""Accuracy cases from a code review (2026-10-16): each was reproduced first, then fixed.

The review's two most serious cases (a name not tied to the business number, or a sister company's name,
confirming a GST/HST number) are in test_vendor_names.py.
"""

import datetime as dt
import tempfile
import unittest
from pathlib import Path

import openpyxl

from taxcheck.checks import Offline, RowChecker, clean_vendor_name, vendor_name_candidates
from taxcheck.numbers import normalize_tax_id
from taxcheck.report import gst_issue, provincial_issue
from taxcheck.result import (NAME_NOT_MATCHED, NOT_APPLICABLE, NOT_LOOKED_UP, NOT_REGISTERED, REGISTERED,
                             Result)
from taxcheck.workbook import SupplierSheet, normalize_country, normalize_province

DATE = dt.date(2026, 9, 1)


class Routing(unittest.TestCase):
    def test_canadian_spellings_are_canadian(self):
        self.assertEqual(normalize_country("CDN"), "CA")
        self.assertEqual(normalize_country("Can."), "CA")
        for value, code in (("Ont.", "ON"), ("Que.", "QC"), ("Man.", "MB"), ("Alta.", "AB"), ("N.S.", "NS"),
                            ("B.C.", "BC"), ("NF", "NL"), ("Newfoundland & Labrador", "NL")):
            self.assertEqual(normalize_province(value), code, value)
        self.assertEqual(normalize_province("CA"), "")  # California, not Canada
        off = Offline()
        rc = RowChecker(off, off, bc=off, mb=off, available={"name", "gst", "country", "province", "pst"})
        out = rc.check({"name": "X LTD", "gst": "108161779", "country": "CDN", "province": "BC", "pst": "1000-7572"}, DATE)
        self.assertEqual(out["gst"].status, NOT_LOOKED_UP)  # checked, not "foreign"
        self.assertIn("bc_pst", out)

    def test_labelled_tax_ids(self):
        for raw in ("GST# 108161779 RT0001", "GST 108161779RT0001", "108161779,RT0001", "108161779RTO001",
                    "GST/HST No. 108161779 RT 0001", "HST: 108161779RT0001"):
            self.assertEqual(normalize_tax_id(raw)[:2], ("gst", "108161779RT0001"), raw)
        self.assertEqual(normalize_tax_id("BN 108161779")[:2], ("gst", "108161779"))

    def test_foreign_tax_id_must_really_be_canadian(self):
        off = Offline()
        rc = RowChecker(off, off, available={"name", "gst", "country"})
        status = lambda tax_id: rc.check({"name": "X CORP", "gst": tax_id, "country": "US"}, DATE)["gst"].status
        self.assertEqual(status("W-8BEN CERT ON FILE"), NOT_APPLICABLE)  # has "RT" in it, but isn't a number
        self.assertEqual(status("PARTNERSHIP"), NOT_APPLICABLE)
        self.assertEqual(status("R108161779"), NOT_LOOKED_UP)  # old-format GST number: checked
        self.assertEqual(status("910000001"), NOT_APPLICABLE)  # bare 9 digits: a US EIN


class Names(unittest.TestCase):
    def test_vendor_master_markers(self):
        cases = {
            "*DNU* ACME LTD": "ACME LTD",
            "ACME LTD (DNU-USE 456)": "ACME LTD",
            "XXX DO NOT USE XXX ACME LTD": "ACME LTD",
        }
        for raw, expected in cases.items():
            self.assertEqual(clean_vendor_name(raw), (expected, True), raw)

    def test_no_junk_candidates(self):
        self.assertEqual(vendor_name_candidates("ACME LTD (O/A ACE)")[-1], "ACE")
        for name, junk in (("R.O.A. CONSULTING LTD", "R."), ("THE OA GROUP INC", "THE"),
                           ("INC RESEARCH LTD", "INC"), ("1234567 ON TIME DELIVERY LTD", "1234567 ONTARIO LTD.")):
            self.assertNotIn(junk, vendor_name_candidates(name), name)
        self.assertIn("1234567 ONTARIO LTD", vendor_name_candidates("1234567 ON LTD"))  # a real numbered company


class Reporting(unittest.TestCase):
    def test_simplified_registrant_with_number_on_file_needs_review(self):
        res = Result(REGISTERED, "", "AIRGSM PTE. LTD.", facts={"simplified_registrant": True, "cra_answer": "registered"})
        issue, action, severity = gst_issue(res)
        self.assertEqual(issue, "Simplified (digital economy) registrant")
        self.assertIn("isn't eligible for input tax credits", action)
        self.assertGreaterEqual(severity, 5)  # Overall REVIEW

    def test_qst_status_read_against_the_payment_date(self):
        cancelled_later = Result(NOT_REGISTERED, "", facts={"registration_end": "2026-10-01"})
        res = RowChecker._qst_on_date(cancelled_later, DATE)
        self.assertEqual(res.status, REGISTERED)
        self.assertEqual(provincial_issue("qst", res)[0], "Registration cancelled since payment date")
        registered_later = Result(REGISTERED, "", facts={"registered_since": "2026-10-01", "registration_end": "Active"})
        res = RowChecker._qst_on_date(registered_later, DATE)
        self.assertEqual(res.status, NOT_REGISTERED)
        self.assertEqual(provincial_issue("qst", res)[0], "Registered after payment date")
        ok = Result(REGISTERED, "", facts={"registered_since": "2017-08-01", "registration_end": "Active"})
        self.assertIs(RowChecker._qst_on_date(ok, DATE), ok)

    def test_trust_gst_number_accepts_any_name_mismatch(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "in.xlsx"
            wb = openpyxl.Workbook()
            wb.active.append(["Supplier Name", "GST/HST Number"])
            wb.active.append(["A CO", "108161779"])
            wb.active.append(["B CO", "857305932"])
            wb.save(path)
            results = {2: {"gst": Result(NAME_NOT_MATCHED, "", "OTHER LTD", facts={"accepted_name": "OTHER LTD"})},
                       3: {"gst": Result(NAME_NOT_MATCHED, "", "", facts={"cancelled_after_payment": True,
                                                                         "registration_end": "2026-08-01"})}}
            SupplierSheet(path).write_results(results, Path(tmp) / "out.xlsx", [], (REGISTERED, NAME_NOT_MATCHED))
            ws = openpyxl.load_workbook(Path(tmp) / "out.xlsx").active
            self.assertEqual(ws["C2"].value, "OK")  # only the name differs: accepted with --trust-gst-number
            self.assertEqual(ws["C3"].value, "REVIEW")  # a cancellation still needs a person


if __name__ == "__main__":
    unittest.main()
