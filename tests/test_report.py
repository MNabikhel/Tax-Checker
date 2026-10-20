"""Report columns: Issue / Action per tax, Main Issue per supplier, the Action List sheet, and the extra
'registered today?' lookup for suppliers that weren't registered on the payment date."""

import datetime as dt
import tempfile
import unittest
from pathlib import Path

import openpyxl

from taxcheck.checks import RowChecker
from taxcheck.report import gst_issue, provincial_issue, row_summary
from taxcheck.result import (
    ERROR, INVALID, MANUAL, MISSING, NAME_NOT_MATCHED, NAME_UNCONFIRMED, NAME_UPDATE, NOT_APPLICABLE, NOT_CONFIRMED, NOT_REGISTERED,
    REGISTERED, Result,
)
from taxcheck.workbook import SupplierSheet


def issue(res):
    return gst_issue(res)[0]


class GstIssues(unittest.TestCase):
    def test_each_situation_has_its_own_issue(self):
        self.assertEqual(issue(Result(REGISTERED, "", facts={})), "Registered - name matches")
        upd = Result(REGISTERED, "", name_check=NAME_UPDATE, facts={"accepted_name": "ACME LTD."})
        self.assertEqual(gst_issue(upd)[:2], ("Registered - vendor name differs from CRA",
                                              "Update the vendor name to 'ACME LTD.' (the name CRA accepted)."))
        rc = Result(REGISTERED, "", facts={"tax_id_problem": "Corporate income tax account (RC), not GST/HST (RT)",
                                           "number_checked": "900000035"})
        self.assertEqual(issue(rc), "Registered - Tax ID needs correcting")
        self.assertIn("900000035RT0001", gst_issue(rc)[1])
        other = Result(NAME_NOT_MATCHED, "", facts={"accepted_name": "ACME INDUSTRIAL CANADA"})
        self.assertEqual(issue(other), "Registered under another name")
        self.assertIn("ACME INDUSTRIAL CANADA", gst_issue(other)[1])
        self.assertEqual(issue(Result(NAME_NOT_MATCHED, "", facts={})), "Registered - legal name not confirmed")
        self.assertEqual(issue(Result(NOT_REGISTERED, "", facts={"cra_answer": "unknown_number"})), "Number not known to CRA")
        late = Result(NOT_REGISTERED, "", facts={"cra_answer": "not_registered_on_date", "status_today": REGISTERED,
                                                 "date_checked": "2026-03-01", "today": "2026-10-13"})
        self.assertEqual(issue(late), "Registered after payment date")
        never = Result(NOT_REGISTERED, "", facts={"cra_answer": "not_registered_on_date", "status_today": NOT_REGISTERED,
                                                  "date_checked": "2026-03-01"})
        self.assertEqual(issue(never), "Not registered")
        self.assertIn("on 2026-03-01 or today", gst_issue(never)[1])
        self.assertEqual(issue(Result(INVALID, "x fails the CRA check digit")), "Tax ID typo (fails check digit)")
        self.assertEqual(issue(Result(INVALID, "not a GST/HST number")), "Tax ID not a GST/HST number")
        self.assertEqual(issue(Result(MISSING, "Tax ID '0' is a placeholder")), "No GST/HST number (placeholder Tax ID)")
        self.assertEqual(issue(Result(MISSING, "No GST/HST number provided.")), "No GST/HST number")
        self.assertEqual(issue(Result(NOT_APPLICABLE, "")), "Foreign supplier - not applicable")
        self.assertEqual(issue(Result(ERROR, "")), "Lookup failed")


class ProvincialIssues(unittest.TestCase):
    def test_provincial_situations(self):
        p = lambda key, res: provincial_issue(key, res)[0]
        self.assertEqual(p("mb_rst", Result(REGISTERED, "active")), "Registered")
        self.assertEqual(p("mb_rst", Result(NOT_CONFIRMED, "found nothing")), "Not found in Manitoba registry")
        self.assertEqual(p("mb_rst", Result(NOT_REGISTERED, "Manitoba RST registry: status 'Inactive' for BN 1.")),
                         "Registration not active (Inactive)")
        self.assertEqual(p("mb_rst", Result(MISSING, "no RST")), "No RST or GST/HST number to search")
        self.assertEqual(p("sk_pst", Result(MANUAL, "captcha")), "Manual check needed (CAPTCHA)")
        self.assertEqual(p("sk_pst", Result(NOT_CONFIRMED, "")), "Not found in Saskatchewan registry")
        self.assertEqual(p("qst", Result(NOT_REGISTERED, "Revenu Québec: registration cancelled since 2020")),
                         "Registration cancelled")
        self.assertEqual(p("qst", Result(REGISTERED, "No TQ suffix was given, so ...")), "Registered (TQ0001 assumed)")
        self.assertEqual(p("bc_pst", Result(NOT_CONFIRMED, "No Match Found")), "PST number doesn't match business number")
        self.assertEqual(p("bc_pst", Result(ERROR, "")), "Lookup failed")

    def test_main_issue_is_the_worst(self):
        taxes = [("gst", "GST/HST"), ("mb_rst", "MB RST"), ("sk_pst", "SK PST")]
        row = {"gst": Result(REGISTERED, "", facts={}), "mb_rst": Result(NOT_CONFIRMED, ""), "sk_pst": Result(MANUAL, "")}
        s = row_summary(row, taxes)
        self.assertEqual(s["main"], "MB RST: Not found in Manitoba registry")
        self.assertEqual(s["other"], "SK PST: Manual check needed (CAPTCHA)")
        self.assertEqual(row_summary({"gst": Result(REGISTERED, "", facts={})}, taxes)["main"], "None")


class RegisteredToday(unittest.TestCase):
    def test_not_registered_on_payment_date_is_checked_again_for_today(self):
        calls = []

        class Gst:
            def check(self, bn9, names, date, extra_names=()):
                calls.append(date)
                if date < dt.date(2026, 6, 1):
                    return Result(NOT_REGISTERED, f"not registered on {date}", facts={"cra_answer": "not_registered_on_date"})
                return Result(REGISTERED, "ok", names[0], facts={"accepted_name": names[0], "name_source": "sheet"})

        checker = RowChecker(Gst(), None, available={"name", "gst"})
        out = checker.check({"name": "NEW CO LTD", "gst": "108161779RT0001"}, dt.date(2026, 3, 1))
        g = out["gst"]
        self.assertEqual(len(calls), 2)
        self.assertEqual(g.facts["status_today"], REGISTERED)
        self.assertEqual(gst_issue(g)[0], "Registered after payment date")
        self.assertEqual(g.facts["date_checked"], "2026-03-01")
        self.assertEqual(g.facts["number_checked"], "108161779RT0001")

    def test_unknown_number_is_not_checked_again(self):
        calls = []

        class Gst:
            def check(self, bn9, names, date, extra_names=()):
                calls.append(date)
                return Result(NOT_REGISTERED, "no account", facts={"cra_answer": "unknown_number"})

        RowChecker(Gst(), None, available={"name", "gst"}).check({"name": "X LTD", "gst": "108161779"}, dt.date(2026, 3, 1))
        self.assertEqual(len(calls), 1)


class Layout(unittest.TestCase):
    def test_columns_action_list_and_summary(self):
        with tempfile.TemporaryDirectory() as tmp:
            path, out = Path(tmp) / "in.xlsx", Path(tmp) / "out.xlsx"
            wb = openpyxl.Workbook()
            wb.active.append(["Vendor ID", "Company Name", "Tax ID", "State"])
            wb.active.append([1, "GOOD CO LTD", "108161779RT0001", "AB"])
            wb.active.append([2, "ACME INC. O/A ACE", "857305932RT0001", "MB"])
            wb.save(path)
            sheet = SupplierSheet(path)
            sheet.write_results({
                2: {"gst": Result(REGISTERED, "ok", facts={"accepted_name": "GOOD CO LTD", "name_source": "sheet",
                                                           "vendor_name": "GOOD CO LTD"})},
                3: {"gst": Result(REGISTERED, "ok", name_check=NAME_UPDATE,
                                  facts={"accepted_name": "ACME INC.", "name_source": "sheet", "vendor_name": "ACME INC. O/A ACE"}),
                    "mb_rst": Result(NOT_CONFIRMED, "found nothing")},
            }, out, ["note"])
            result = openpyxl.load_workbook(out)
            self.assertEqual(result.sheetnames, ["Sheet", "Action List", "Tax Check Summary"])
            ws = result["Sheet"]
            headers = [c.value for c in ws[1]]
            for h in ("Main Issue", "Action Needed", "Suggested Vendor Name", "GST/HST Issue", "GST/HST Tax ID Problem",
                      "GST/HST Name Accepted by CRA", "GST/HST Name Source", "MB RST Issue", "QST Issue"):
                self.assertIn(h, headers)
            row = dict(zip(headers, [c.value for c in ws[3]]))
            self.assertEqual(row["Main Issue"], "MB RST: Not found in Manitoba registry")
            self.assertEqual(row["Other Issues"], "GST/HST: Registered - vendor name differs from CRA")
            self.assertEqual(row["Suggested Vendor Name"], "ACME INC.")
            self.assertEqual(row["GST/HST Name Source"], "Vendor name, reworded (legal part / suffix)")
            actions = list(result["Action List"].iter_rows(values_only=True))
            self.assertEqual(len(actions), 2)  # header + ACME (GOOD CO needs nothing)
            self.assertEqual(actions[1][0], 2)
            self.assertEqual(actions[1][1], "ACME INC. O/A ACE")


class CrossRowNotes(unittest.TestCase):
    """Findings that need the whole sheet: a GST/HST number on differently named vendors, and cut-off names."""

    def write(self, rows, results):
        with tempfile.TemporaryDirectory() as tmp:
            path, out = Path(tmp) / "in.xlsx", Path(tmp) / "out.xlsx"
            wb = openpyxl.Workbook()
            wb.active.append(["Vendor ID", "Company Name", "Tax ID", "State"])
            for row in rows:
                wb.active.append(row)
            wb.save(path)
            SupplierSheet(path).write_results(results, out, [])
            result = openpyxl.load_workbook(out)
            ws = result["Sheet"]
            headers = [c.value for c in ws[1]]
            table = {r: dict(zip(headers, [c.value for c in ws[r]])) for r in range(2, ws.max_row + 1)}
            summary = [r for r in result["Tax Check Summary"].iter_rows(values_only=True)]
            return table, summary

    def test_number_shared_by_differently_named_vendors(self):
        number = {"number_checked": "840098545RT0001"}
        table, summary = self.write(
            [[1, "(DNU) WILLOW FLORIST", "840098545RT0001", "MB"], [2, "(DNU) GIFT-BOX", "840098545RT0001", "MB"],
             [3, "ACME LTD", "857305932RT0001", "MB"], [4, "ACME LTD.", "857305932RT0001", "MB"]],
            {2: {"gst": Result(REGISTERED, "ok", "WILLOW FLORIST", facts={**number, "accepted_name": "WILLOW FLORIST"})},
             3: {"gst": Result(NAME_NOT_MATCHED, "not under the name", facts=dict(number))},
             4: {"gst": Result(REGISTERED, "ok", facts={"number_checked": "857305932RT0001"})},
             5: {"gst": Result(REGISTERED, "ok", facts={"number_checked": "857305932RT0001"})}})
        self.assertIn("same number as row 3 '(DNU) GIFT-BOX'", table[2]["Other Issues"])
        self.assertIn("row 2 '(DNU) WILLOW FLORIST' (CRA accepts the number under 'WILLOW FLORIST')",
                      table[3]["Other Issues"])
        self.assertIsNone(table[4]["Other Issues"])  # the same vendor twice, spelt differently: not flagged
        self.assertTrue(any(r[1] == "Same GST/HST number as another vendor" and r[2] == 2 for r in summary))

    def test_same_vendor_spelt_differently_or_invalid_numbers_arent_flagged(self):
        from taxcheck.result import INVALID

        table, _ = self.write(
            [[1, "ACME", "123456782RT0001", "AB"], [2, "ACME INC.", "123456782RT0001", "AB"],
             [3, "Café Ltée", "857305932RT0001", "QC"], [4, "CAFE LTEE", "857305932RT0001", "QC"],
             [5, "FIRST CO", "111111111", "AB"], [6, "SECOND CO", "111111111", "AB"]],
            {2: {"gst": Result(REGISTERED, "ok", facts={"number_checked": "123456782RT0001"})},
             3: {"gst": Result(REGISTERED, "ok", facts={"number_checked": "123456782RT0001"})},
             4: {"gst": Result(REGISTERED, "ok", facts={"number_checked": "857305932RT0001"})},
             5: {"gst": Result(REGISTERED, "ok", facts={"number_checked": "857305932RT0001"})},
             6: {"gst": Result(INVALID, "check digit", facts={"number_checked": "111111111"})},
             7: {"gst": Result(INVALID, "check digit", facts={"number_checked": "111111111"})}})
        for r in range(2, 8):
            self.assertFalse(table[r]["Other Issues"], r)

    def test_name_at_the_vendor_masters_length_limit(self):
        long_names = ["ANCHOR CONSTRUCTION INDUSTRIAL PRODUCT", "NORTHERN PRAIRIE MECHANICAL CONTRACTOR",
                      "WESTERN INDUSTRIAL SUPPLY AND SERVICES"]
        self.assertEqual({len(n) for n in long_names}, {38})
        unconfirmed = Result(NAME_NOT_MATCHED, "not under the name", name_check=NAME_UNCONFIRMED)
        table, _ = self.write([[i, n, "", "AB"] for i, n in enumerate(long_names + ["SHORT NAME LTD"])],
                              {2: {"gst": unconfirmed}, 3: {"gst": Result(REGISTERED, "ok")},
                               4: {"gst": unconfirmed.copy()}, 5: {"gst": unconfirmed.copy()}})
        self.assertIn("probably cut off", table[2]["Action Needed"])
        self.assertIn("probably cut off", table[4]["Action Needed"])
        self.assertNotIn("cut off", table[5]["Action Needed"] or "")  # shorter than the limit


class SaskatchewanLabels(unittest.TestCase):
    def test_manual_and_near_match_reasons(self):
        self.assertEqual(provincial_issue("sk_pst", Result(MANUAL, "needs at least 4 letters of the name"))[0],
                         "Manual check needed (name too short to search)")
        self.assertIn("daily search limit", provincial_issue("sk_pst", Result(
            MANUAL, "allows 100 searches per 24 hours and they're used up"))[0])
        self.assertEqual(provincial_issue("sk_pst", Result(NOT_CONFIRMED, "No result is exactly the supplier's name; "
                                                           "similar names in the registry: X"))[0],
                         "Only similar names in Saskatchewan registry")
        self.assertEqual(provincial_issue("sk_pst", Result(NOT_CONFIRMED, "found no business"))[0],
                         "Not found in Saskatchewan registry")


if __name__ == "__main__":
    unittest.main()
