"""Tailoring to the ERP vendor export (made-up rows in the same layout as the real file).

The real export (vendor_export.xlsx, ~20,000 vendors, Canadian and foreign) has one free-form
"Tax ID" column, a Country column, State, Postal Code and "Check/ Itm Date" (last payment), and no QST or
PST columns. These tests pin down how that file is read and routed, plus the progress/resume features
added for runs that take hours.
"""

import csv
import datetime as dt
import logging
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import openpyxl

import check_suppliers
from taxcheck.checks import RowChecker
from taxcheck.numbers import normalize_tax_id
from taxcheck.progress import load_resume
from taxcheck.result import ERROR, MISSING, NOT_APPLICABLE, REGISTERED, Result
from taxcheck.workbook import SupplierSheet, province_from_postal

HEADERS = [
    "Internal Supplier Number", "Company Name", "Complete Address", "Address Line 1", "Address Line 2",
    "Address Line 3", "Address Line 4", "City - CTY1", "State", "Country", "Postal Code", "Tax ID", "Cur Cod",
    "Check/ Itm Date", "Email Address", "Payment Amount", "Total Payment Count", "Pmt Trm",
    "Payment Terms - A/P", "Business Unit", "Business Unit", "DUNS Number", "Sch Typ",
]


def vendor(num, name, state, country, postal, tax_id, date=dt.datetime(2026, 9, 1)):
    return [num, f" {name}", f" {name}\n1 Main St\n", "1 Main St", "", "", "", "Town", state, country, postal,
            tax_id, "CAD", date, "", -100.0, 1, "A", "Net 30", "  0659999999", "ADMIN", "", "V"]


ROWS = [
    vendor(1, "MAPLE BUILDERS LTD", "AB", "CA", "T5J 0N3", "108161779RT0001"),          # 2 plain GST
    vendor(2, "9999-0000 QUEBEC INC", "QC", "CA", "J6J 4Z2", "108161779RC0001"),        # 3 RC account
    vendor(3, "ACME GLASS CO. INC.", "LA", "US", "71201", "910000001"),                 # 4 US EIN -> N/A
    vendor(4, "GLOBEX, INC.", "IL", "US", "60601", "857305932RT0001"),                  # 5 US with GST
    vendor(5, "JANE DOE", "ON", "CA", "M3H 0A1", "000000000RT0001"),                   # 6 placeholder
    vendor(6, "PRAIRIE SUPPLY", "CA", "CA", "R3C 1A1", "R108161779"),                  # 7 state "CA", old R
    vendor(7, "BAYSIDE CONSTRUCTION", "WA", None, "98206", "920000002"),             # 8 blank country, US
    vendor(8, "SERVICES ABC INC", "QC", "CA", "H2X 1Y4", "1019288451TQ0004"),          # 9 QST in Tax ID
    vendor(9, "COAST TOOLS LTD", "BC", "CA", "V6C 3A6", "108161779RT0001"),            # 10 BC, no PST column
]


class RecordingGst:
    def __init__(self, *a, **k):
        self.simplified = mock.Mock(error=None)
        self.calls = []

    def check(self, bn9, names, date, extra_names=()):
        self.calls.append((bn9, date))
        return Result(REGISTERED, "fake CRA", names[0])


class FakeQst:
    nr_error = None

    def __init__(self, *a, **k):
        self.calls = []

    def check(self, number, suffix_assumed=False):
        self.calls.append(number)
        return Result(REGISTERED, "fake RQ")


class FakeMb:
    def __init__(self):
        self.calls = []

    def check(self, names, bn9=None, rst=None):
        self.calls.append((bn9, rst))
        return Result(REGISTERED, "fake TAXcess")


def make_export_with(path, headers, rows):
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet1"
    wb.active.append(headers)
    for r in rows:
        wb.active.append(r)
    wb.save(path)


def make_export(path, rows=ROWS):
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet1"
    wb.active.append(HEADERS)
    for r in rows:
        wb.active.append(r)
    wb.save(path)


class TaxIdNormalization(unittest.TestCase):
    def test_forms_seen_in_the_export(self):
        self.assertEqual(normalize_tax_id("108161779RT0001")[:2], ("gst", "108161779RT0001"))
        self.assertEqual(normalize_tax_id("108161779 RT0001")[:2], ("gst", "108161779RT0001"))
        kind, num, note = normalize_tax_id("108161779RC0001")
        self.assertEqual((kind, num), ("gst", "108161779"))
        self.assertIn("corporate income tax", note)
        self.assertEqual(normalize_tax_id("R108161779")[:2], ("gst", "108161779"))
        self.assertEqual(normalize_tax_id("R-108161779")[:2], ("gst", "108161779"))
        self.assertEqual(normalize_tax_id("108161779RT001")[:2], ("gst", "108161779"))
        self.assertIn("typo", normalize_tax_id("108161779TR0001")[2])
        self.assertEqual(normalize_tax_id("000000000RT0001")[0], "zero")
        self.assertEqual(normalize_tax_id("00000000RT00001")[0], "zero")
        self.assertEqual(normalize_tax_id("1019288451TQ0004")[:2], ("qst", "1019288451TQ0004"))
        self.assertEqual(normalize_tax_id("1234567RT0001")[0], "raw")   # 7 digits: left for parse_gst to reject
        self.assertEqual(normalize_tax_id("12 345 678 901")[0], "raw")  # Australian ABN

    def test_do_not_use_markers_removed_from_lookup_name(self):
        from taxcheck.checks import clean_vendor_name

        cases = {
            "(DNU) SUMMIT CONSTRUCTION INDUSTRIAL PRO": "SUMMIT CONSTRUCTION INDUSTRIAL PRO",
            "DNU - BEACON SECURITY SERVICES": "BEACON SECURITY SERVICES",
            "DO NOT USE-STONEFIELD CONCRETE LTD": "STONEFIELD CONCRETE LTD",
            "PEAKLINE CONSTRUCTION INC. DO NOT USE": "PEAKLINE CONSTRUCTION INC.",
            "NORTHFIELD SUPPLY INC -REMIT DO NOT USE": "NORTHFIELD SUPPLY INC",
            "TRENCHLINE CANADA INC. (USE 107 DO NOT USE": "TRENCHLINE CANADA INC.",
            "TRADES LOCAL 999-UNION PROM(DNU)": "TRADES LOCAL 999-UNION PROM",
            "TELEMONDE  LTEE DNU": "TELEMONDE LTEE",
        }
        for raw, expected in cases.items():
            self.assertEqual(clean_vendor_name(raw), (expected, True), raw)
        for untouched in ("DNUTECH LTD", "SUNDANCE INC", "ACME  SUPPLY CO.   1840", None):
            self.assertEqual(clean_vendor_name(untouched), (untouched, False))
        gst = RecordingGst()
        checker = RowChecker(gst, FakeQst(), available={"name", "gst"})
        out = checker.check({"name": "(DNU) MAPLE BUILDERS LTD", "gst": "108161779RT0001"}, dt.date(2026, 9, 1))
        self.assertIn("marked do-not-use", out["gst"].detail)
        self.assertEqual(out["gst"].registered_name, "MAPLE BUILDERS LTD")  # the cleaned name was sent

    def test_postal_code_province(self):
        self.assertEqual(province_from_postal("T5J 0N3"), "AB")
        self.assertEqual(province_from_postal("v6c3a6"), "BC")
        self.assertEqual(province_from_postal("98206"), "")
        self.assertEqual(province_from_postal(None), "")


class ExportReading(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "vendor_export.xlsx"
        make_export(self.path)

    def tearDown(self):
        self.tmp.cleanup()

    def test_columns_mapped(self):
        sheet = SupplierSheet(self.path)
        names = {k: sheet.ws.cell(1, c).value for k, c in sheet.cols.items()}
        self.assertEqual(names, {"name": "Company Name", "province": "State", "country": "Country",
                                 "postal": "Postal Code", "gst": "Tax ID", "date": "Check/ Itm Date"})
        (r, first), *_ = sheet.rows()
        self.assertEqual(first["name"], "MAPLE BUILDERS LTD")  # leading space from the export is trimmed

    def route(self):
        sheet = SupplierSheet(self.path)
        gst, qst, mb = RecordingGst(), FakeQst(), FakeMb()
        checker = RowChecker(gst, qst, bc=mock.Mock(), mb=mb, sk=None, available=sheet.cols)
        out = {r: checker.check(v, dt.date(2026, 9, 1)) for r, v in sheet.rows()}
        return out, gst, qst, mb

    def test_payment_amount_column_not_needed(self):
        """The user deletes "Payment Amount" from the export: nothing reads it, and columns are found by name."""
        trimmed = Path(self.tmp.name) / "trimmed.xlsx"
        i = HEADERS.index("Payment Amount")
        make_export_with(trimmed, [h for j, h in enumerate(HEADERS) if j != i], [[v for j, v in enumerate(r) if j != i] for r in ROWS])
        full, cut = SupplierSheet(self.path), SupplierSheet(trimmed)
        self.assertEqual(set(full.cols), set(cut.cols))
        self.assertEqual([v for _, v in full.rows()], [v for _, v in cut.rows()])

    def test_routing(self):
        out, gst, qst, mb = self.route()
        self.assertEqual(out[2]["gst"].status, REGISTERED)
        self.assertEqual(set(out[2]), {"gst"})
        self.assertIn("corporate income tax", out[3]["gst"].detail)
        self.assertNotIn("qst", out[3])  # QC, but the export has no QST column: not flagged NO NUMBER
        self.assertEqual(out[4], {"gst": out[4]["gst"]})
        self.assertEqual(out[4]["gst"].status, NOT_APPLICABLE)
        self.assertEqual(out[5]["gst"].status, REGISTERED)  # foreign supplier with an explicit GST number
        self.assertIn("Non-Canadian supplier (country US)", out[5]["gst"].detail)
        self.assertEqual(out[6]["gst"].status, MISSING)
        self.assertIn("placeholder", out[6]["gst"].detail)
        self.assertIn("mb_rst", out[7])  # State "CA" -> Manitoba from postal code R3C
        self.assertEqual(mb.calls, [("108161779", None)])
        self.assertEqual(out[8]["gst"].status, NOT_APPLICABLE)  # blank country, US state
        self.assertEqual(qst.calls, ["1019288451TQ0004"])
        self.assertEqual(out[9]["qst"].status, REGISTERED)
        self.assertEqual(out[9]["gst"].status, MISSING)
        self.assertNotIn("bc_pst", out[10])  # no PST column in the export
        self.assertNotIn(("910000001", dt.date(2026, 9, 1)), gst.calls)  # US EIN never sent to CRA


class ExportCli(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.src = self.dir / "vendor_export.xlsx"
        make_export(self.src)
        self.logs = self.dir / "logs"
        for p in (mock.patch.object(check_suppliers, "GstChecker", RecordingGst),
                  mock.patch.object(check_suppliers, "QstChecker", FakeQst)):
            p.start()
            self.addCleanup(p.stop)

    def tearDown(self):
        root = logging.getLogger()
        for h in list(root.handlers):
            root.removeHandler(h)
            h.close()
        self.tmp.cleanup()

    def run_cli(self, *extra):
        if "--workers" not in extra:
            extra = ("--workers", "1", *extra)  # deterministic order for these tests; parallel has its own test
        return check_suppliers.main([str(self.src), "--log-dir", str(self.logs), "--no-browser", "--no-name-lookup",
                                     *extra])

    def results(self):
        ws = openpyxl.load_workbook(self.dir / "vendor_export_tax_check.xlsx")["Sheet1"]
        headers = [c.value for c in ws[1]]
        return [dict(zip(headers, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]

    def test_offline_run_writes_every_row_and_progress_file(self):
        self.assertEqual(self.run_cli("--offline"), 0)
        rows = self.results()
        self.assertEqual(len(rows), len(ROWS))
        overall = [r["Overall"] for r in rows]
        self.assertEqual(overall[0], "PENDING")
        self.assertEqual(overall[2], "NOT APPLICABLE")
        self.assertEqual(overall[4], "REVIEW")  # placeholder Tax ID
        ws = openpyxl.load_workbook(self.dir / "vendor_export_tax_check.xlsx")["Sheet1"]
        self.assertEqual(ws["T2"].value, "  0659999999")  # original columns untouched (first "Business Unit")
        self.assertEqual(ws["X1"].value, "Overall")  # results start right after the export's last column
        progress = list(self.logs.glob("progress_*.csv"))
        self.assertEqual(len(progress), 1)
        with open(progress[0], newline="", encoding="utf-8-sig") as f:
            lines = list(csv.DictReader(f))
        self.assertEqual(len(lines), len(ROWS))
        self.assertEqual(lines[0]["province"], "AB")
        self.assertIn("PROGRESS 9/9", next(self.logs.glob("tax_check_*.log")).read_text(encoding="utf-8"))

    def test_limit_then_resume_finishes_the_rest_and_retries_errors(self):
        self.assertEqual(self.run_cli("--limit", "3"), 0)
        self.assertTrue((self.dir / "vendor_export_tax_check_TRIAL.xlsx").exists())  # trial has its own file
        self.assertFalse((self.dir / "vendor_export_tax_check.xlsx").exists())
        first = next(self.logs.glob("progress_*.csv"))
        done = load_resume(first)
        # 3 lookups: rows 2 and 3 (Canadian), row 4 (US, N/A, passed on the way), row 5 (US with a GST number)
        self.assertEqual(sorted(done), [2, 3, 4, 5])
        # Pretend row 3 had failed: it must be checked again on resume.
        with open(first, newline="", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            fields, lines = reader.fieldnames, list(reader)
        lines[1]["results_json"] = '{"gst": ["ERROR", "simulated", ""]}'
        with open(first, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=fields)
            writer.writeheader()
            writer.writerows(lines)
        self.assertEqual(sorted(load_resume(first)), [2, 4, 5])
        with mock.patch("taxcheck.progress.dt") as fake_dt:  # a different progress file name for the second run
            fake_dt.datetime.now.return_value = dt.datetime(2030, 1, 1)
            self.assertEqual(self.run_cli("--resume", str(first)), 0)
        rows = self.results()
        self.assertTrue(all(r["Overall"] for r in rows))
        self.assertEqual(rows[1]["GST/HST Status"], REGISTERED)
        latest = self.logs / "progress_20300101_000000.csv"
        with open(latest, newline="", encoding="utf-8-sig") as f:
            sources = [line["source"] for line in csv.DictReader(f)]
        self.assertEqual(sources.count("resumed"), 3)
        self.assertEqual(sources.count("checked"), len(ROWS) - 3)

    def test_checkpoint_saves_during_the_run(self):
        saves = []
        real = check_suppliers.save_workbook

        def spy(sheet, results, out, notes, ok):
            saves.append((len(results), notes[0]))
            return real(sheet, results, out, notes, ok)

        with mock.patch.object(check_suppliers, "save_workbook", spy):
            self.assertEqual(self.run_cli("--offline", "--checkpoint-every", "4"), 0)
        self.assertEqual([n for n, _ in saves], [4, 8, 9])
        self.assertTrue(saves[0][1].startswith("IN PROGRESS"))

    def test_parallel_run_matches_one_at_a_time_and_browser_checks_stay_on_main_thread(self):
        import threading
        import time as _time

        many = [vendor(100 + i, f"VENDOR {i} LTD", "MB" if i % 5 == 0 else "AB", "CA",
                       "R3C 1A1" if i % 5 == 0 else "T5J 0N3", "108161779RT0001" if i % 2 else "857305932RT0001",
                       dt.datetime(2026, 9, 1 + i % 20)) for i in range(40)]
        make_export(self.src, ROWS + many)
        threads, mb_threads = set(), []

        class SlowGst(RecordingGst):
            def check(self_, bn9, names, date, extra_names=()):
                threads.add(threading.current_thread().name)
                _time.sleep(0.05)
                return super().check(bn9, names, date, extra_names)

        class MainOnlyMb:
            def __init__(self, browser):
                pass

            def check(self_, names, bn9=None, rst=None):
                mb_threads.append(threading.current_thread() is threading.main_thread())
                return Result(REGISTERED, "fake TAXcess")

        fake_browser = mock.MagicMock()
        fake_browser.__enter__.return_value = fake_browser
        with mock.patch.object(check_suppliers, "GstChecker", SlowGst), \
                mock.patch.object(check_suppliers, "MbRstChecker", MainOnlyMb), \
                mock.patch.object(check_suppliers, "Browser", return_value=fake_browser):
            t0 = _time.monotonic()
            self.assertEqual(check_suppliers.main([str(self.src), "--log-dir", str(self.logs), "--no-name-lookup",
                                                   "--workers", "4", "--as-of-today"]), 0)
            parallel_secs = _time.monotonic() - t0
        rows = self.results()
        self.assertEqual(len(rows), len(ROWS) + 40)
        self.assertTrue(all(r["Overall"] for r in rows))
        self.assertGreater(len(threads), 1)  # GST lookups really ran on several threads
        self.assertTrue(mb_threads and all(mb_threads))  # every Manitoba check ran on the main thread
        self.assertTrue(all(r["MB RST Status"] == REGISTERED for r in rows if r["State"] == "MB"))
        self.assertEqual([r["Company Name"].strip() for r in rows][-1], "VENDOR 39 LTD")  # sheet order kept
        log_text = next(self.logs.glob("tax_check_*.log")).read_text(encoding="utf-8")
        self.assertIn("Checking with 4 parallel lookups", log_text)
        self.assertLess(parallel_secs, 40 * 0.05)  # faster than one at a time

    def test_gst_only_run_skips_provincial_checks_and_columns(self):
        mb_calls = []

        class Mb:
            def __init__(self, browser):
                pass

            def check(self_, *a, **k):
                mb_calls.append(a)
                return Result(REGISTERED, "fake")

        with mock.patch.object(check_suppliers, "MbRstChecker", Mb), \
                mock.patch.object(check_suppliers, "Browser", return_value=mock.MagicMock()):
            self.assertEqual(check_suppliers.main([str(self.src), "--log-dir", str(self.logs), "--no-name-lookup",
                                                   "--workers", "1", "--taxes", "gst"]), 0)
        self.assertEqual(mb_calls, [])
        wb = openpyxl.load_workbook(self.dir / "vendor_export_tax_check.xlsx")
        headers = [c.value for c in wb["Sheet1"][1]]
        self.assertIn("GST/HST Status", headers)
        self.assertFalse(any(h and ("MB RST" in h or "QST" in h or "SK PST" in h) for h in headers))
        notes = [r[0] for r in wb["Tax Check Summary"].iter_rows(values_only=True) if r and r[0]]
        self.assertIn("Taxes checked in this run: GST/HST.", notes)

    def test_tax_choice_parsing_and_menu(self):
        self.assertIsNone(check_suppliers.parse_taxes("all"))
        self.assertEqual(check_suppliers.parse_taxes("gst, mb"), {"gst", "mb_rst"})
        self.assertEqual(check_suppliers.parse_taxes("M"), {"gst", "mb_rst"})  # GST/HST always included
        self.assertEqual(check_suppliers.parse_taxes("G Q B M S"), {"gst", "qst", "bc_pst", "mb_rst", "sk_pst"})
        with self.assertRaises(ValueError):
            check_suppliers.parse_taxes("vat")
        for answers, expected in ((["1"], {"gst"}), ([""], None), (["2"], None), (["3", "g m"], {"gst", "mb_rst"}),
                                  (["x", "1"], {"gst"})):
            with mock.patch("builtins.input", side_effect=answers), mock.patch("builtins.print"):
                self.assertEqual(check_suppliers.ask_taxes(), expected, answers)

    def test_unexpected_failure_keeps_results_so_far(self):
        calls = []

        def boom(self_, row, date):
            calls.append(row)
            if len(calls) == 3:
                raise MemoryError("simulated")
            return {"gst": Result(REGISTERED, "ok")}

        with mock.patch.object(RowChecker, "check", boom):
            self.assertEqual(self.run_cli(), 1)
        rows = self.results()
        self.assertEqual([r["Overall"] for r in rows[:3]], ["OK", "OK", None])
        log_text = next(self.logs.glob("tax_check_*.log")).read_text(encoding="utf-8")
        self.assertIn("Run stopped by an unexpected error after 2 suppliers", log_text)


if __name__ == "__main__":
    unittest.main()
