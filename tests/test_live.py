"""Live checks against the real government services, using publicly published registrations.

Skipped unless TAXCHECK_LIVE=1, because they need internet access and take a minute or two:

    TAXCHECK_LIVE=1 python -m unittest tests.test_live -v

The Saskatchewan test also needs a person: with TAXCHECK_SK_ASSIST=1 it opens a browser window where
you tick the registry's "I'm not a robot" box, then searches automatically (up to 5 minutes to tick it).

Run these when something looks off: a failure here usually means a government site changed.
"""

import datetime as dt
import os
import unittest

from taxcheck.certs import use_system_certificates
from taxcheck.gst import GstChecker, cra_today
from taxcheck.qst import QstChecker
from taxcheck.result import INVALID, NOT_CONFIRMED, NOT_REGISTERED, REGISTERED

LIVE = os.environ.get("TAXCHECK_LIVE") == "1"
if LIVE:
    use_system_certificates()  # corporate networks re-sign HTTPS traffic; trust Windows' certificates


def reachable(url):
    """True if the site answers at all, so an outage is reported as a skip rather than a code failure."""
    import requests

    try:
        requests.get(url, timeout=20)
        return True
    except requests.RequestException:
        return False
TODAY = cra_today()


@unittest.skipUnless(LIVE, "set TAXCHECK_LIVE=1 to run live tests")
class LiveCra(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.gst = GstChecker(delay=1.0)

    def test_registry_outcomes(self):
        lookup = self.gst.registry_lookup
        self.assertEqual(lookup("857305932", "Amazon.com.ca ULC", TODAY), "registered")
        self.assertEqual(lookup("857305932", "Wrong Name Co", TODAY), "no_match")
        self.assertEqual(lookup("857305932", "Amazon.com.ca ULC", dt.date(1990, 1, 1)), "not_registered")
        # CRA answers "not registered on this date" before it looks at the name.
        self.assertEqual(lookup("857305932", "Wrong Name Co", dt.date(1990, 1, 1)), "not_registered")

    def test_full_check_with_name_variant(self):
        res = self.gst.check("108161779", ["The University of British Columbia"], TODAY)
        self.assertEqual(res.status, REGISTERED, res.detail)

    def test_simplified_registrant(self):
        res = self.gst.check("751950577", ["AIRGSM PTE. LTD."], TODAY)
        self.assertEqual(res.status, REGISTERED, res.detail)
        self.assertIsNone(self.gst.simplified.error)
        self.assertGreater(len(self.gst.simplified._rows), 1000)
        self.assertIn("input tax credits", res.detail)

    def test_number_only(self):
        from taxcheck.result import NAME_NOT_MATCHED

        g = GstChecker(delay=1.0)  # no federal names: the number alone decides
        self.assertEqual(g.check("857305932", [None], TODAY).status, NAME_NOT_MATCHED)  # no name in the sheet
        self.assertEqual(g.check("857305932", ["Totally Wrong Vendor"], TODAY).status, NAME_NOT_MATCHED)
        self.assertEqual(g.check("723754966", ["Wrong"], TODAY).status, NOT_REGISTERED)  # existing BN, not registered
        res = g.check("298277534", ["Wrong"], TODAY)  # check-digit-valid number CRA doesn't know
        self.assertEqual(res.status, NOT_REGISTERED, res.detail)
        self.assertIn("no GST/HST account", res.detail)

    def test_not_registered_before_registration(self):
        res = self.gst.check("857305932", ["Amazon.com.ca ULC"], dt.date(1990, 1, 1))
        self.assertEqual(res.status, NOT_REGISTERED, res.detail)


@unittest.skipUnless(LIVE, "set TAXCHECK_LIVE=1 to run live tests")
class LiveQst(unittest.TestCase):
    def test_api(self):
        qst = QstChecker()
        res = qst.check("1019288451TQ0004")
        self.assertEqual(res.status, REGISTERED, res.detail)
        self.assertIn("AVIENT", res.registered_name)
        self.assertEqual(qst.check("1234567890TQ0001").status, INVALID)
        self.assertEqual(qst.check("1019288451TQ0001").status, NOT_REGISTERED)  # cancelled account


@unittest.skipUnless(LIVE, "set TAXCHECK_LIVE=1 to run live tests")
class LiveBrowserChecks(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from taxcheck.browser import Browser

        cls.browser = Browser().__enter__()

    @classmethod
    def tearDownClass(cls):
        cls.browser.__exit__(None, None, None)

    def test_etaxbc(self):
        from taxcheck.bc_pst import ETAXBC, BcPstChecker

        if not reachable(ETAXBC):
            self.skipTest("eTaxBC isn't responding (it has scheduled maintenance windows); try again later")

        bc = BcPstChecker(self.browser)
        self.assertEqual(bc.check("108161779", "1000-7572").status, REGISTERED)
        self.assertEqual(bc.check("108161779", "1000-7573").status, NOT_CONFIRMED)

    def test_manitoba_taxcess(self):
        from taxcheck.mb_rst import MbRstChecker

        mb = MbRstChecker(self.browser)
        res = mb.check(["Amazon Canada Fulfillment"], bn9="857305932")  # falls back to the first word
        self.assertEqual(res.status, REGISTERED, res.detail)
        self.assertIn("AMAZON.COM.CA ULC", res.registered_name)
        self.assertEqual(mb.check(["University of British Columbia"], bn9="108161779").status, NOT_CONFIRMED)
        # A name with an apostrophe finds nothing on TAXcess; the apostrophe-free part does.
        self.assertEqual(mb.check(["Amazon's"], bn9="857305932").status, REGISTERED)

    def test_simplified_list_through_the_browser(self):
        # canada.ca stalls Python on some corporate networks; the browser copy must parse the same.
        import requests
        from unittest import mock

        from taxcheck.gst import SimplifiedList

        with mock.patch("taxcheck.gst.requests.get", side_effect=requests.exceptions.ReadTimeout("simulated stall")):
            sl = SimplifiedList(requests.Session(), browser=self.browser)
            entry = sl.lookup("751950577", TODAY)
        self.assertIsNone(sl.error)
        self.assertEqual(entry["legal_name"], "AIRGSM PTE. LTD.")
        self.assertGreater(len(sl._rows), 500)


@unittest.skipUnless(LIVE and os.name == "nt", "Windows only: Microsoft Edge through Playwright")
class LiveEdge(unittest.TestCase):
    def test_edge_starts(self):
        from playwright.sync_api import sync_playwright

        with sync_playwright() as pw:
            browser = pw.chromium.launch(channel="msedge", headless=True)
            try:
                page = browser.new_page()
                self.assertLess(page.goto("https://taxcess.gov.mb.ca/TAXcess/?Link=RSTLookup").status, 400)
            finally:
                browser.close()


@unittest.skipUnless(LIVE, "set TAXCHECK_LIVE=1 to run live tests")
class LiveQuebecNrList(unittest.TestCase):
    def test_nr_list_loads_and_checks(self):
        from taxcheck.browser import Browser

        with Browser() as browser:
            qst = QstChecker(browser=browser)  # no cache: always the live page
            rows = qst._load_nr_list()
        # Skip only if every source was refused outright (Revenu Québec blocks some networks). A page that
        # loads but yields no NR numbers is a real failure: the parser needs updating.
        if not rows and qst.nr_refusals and "had no NR numbers" not in (qst.nr_error or ""):
            self.skipTest(f"Revenu Québec refused this network: {qst.nr_error}")
        self.assertIsNone(qst.nr_error)
        self.assertGreater(len(rows), 1000, "the NR list should have ~2,000 registrants")
        number, entry = next(iter(rows.items()))
        self.assertTrue(entry["legal_name"] or entry["trade_name"], f"no name parsed for {number}: {entry}")
        self.assertEqual(qst.check_nr(number).status, REGISTERED)
        print(f"\n  NR list: {len(rows)} registrants; first: {number} {entry}")


@unittest.skipUnless(LIVE and os.environ.get("TAXCHECK_SK_ASSIST") == "1",
                     "set TAXCHECK_LIVE=1 and TAXCHECK_SK_ASSIST=1 (needs a person to tick a CAPTCHA)")
class LiveSaskatchewanAssisted(unittest.TestCase):
    def test_known_vendor(self):
        from taxcheck.browser import Browser
        from taxcheck.sk_pst import SkPstAssistant

        with Browser(headless=False, debug_dir=".") as browser:
            sk = SkPstAssistant(browser)
            # Federated Co-operatives Limited (Saskatoon) sells at retail across Saskatchewan.
            res = sk.check(["Federated Co-operatives Limited"])
        print(f"\n  SK result: {res.status} | {res.detail} | {res.registered_name}")
        self.assertEqual(res.status, REGISTERED, "check the sk_*.png screenshot and the log, then adjust sk_pst.classify()")


@unittest.skipUnless(LIVE, "set TAXCHECK_LIVE=1 to run live tests")
class LiveOrgBook(unittest.TestCase):
    def test_bc_company_name_recovered(self):
        from taxcheck.orgbook import OrgBookBC
        from taxcheck.result import NAME_NOT_MATCHED

        ob = OrgBookBC()
        self.assertIn("2K PLUMBING LTD.", ob.names("859512196"))  # BC-incorporated, not in the federal data
        res = GstChecker(delay=1.0, name_sources=[ob]).check("859512196", ["2K Plumbing & Heating"], TODAY)
        self.assertEqual(res.status, NAME_NOT_MATCHED, res.detail)
        self.assertEqual(res.registered_name, "2K PLUMBING LTD.")
        self.assertIn("OrgBook BC", res.detail)


@unittest.skipUnless(LIVE, "set TAXCHECK_LIVE=1 to run live tests")
class LiveFederalNames(unittest.TestCase):
    def test_name_mismatch_recovered(self):
        from taxcheck.fedcorp import FederalCorporations

        gst = GstChecker(delay=1.0, corporations=FederalCorporations(cache_dir="cache"))
        res = gst.check("774075766", ["NA Tutors"], TODAY)  # sheet name doesn't match CRA's
        from taxcheck.result import NAME_NOT_MATCHED

        self.assertEqual(res.status, NAME_NOT_MATCHED, res.detail)
        self.assertIn("North American Tutors", res.detail)


@unittest.skipUnless(LIVE, "set TAXCHECK_LIVE=1 to run live tests")
class LiveCraNameRules(unittest.TestCase):
    """How CRA compares names (gst._name_key relies on it). Measured 2026-10-16."""

    def test_spacing_and_punctuation(self):
        from taxcheck.gst import _name_key

        g = GstChecker(delay=1.0)
        cases = [("857305932", "Amazon com ca ULC"), ("857305932", "Amazoncomca ULC"),
                 ("857305932", "Amazon&com&ca ULC"), ("857305932", "Amazon,com,ca ULC"),
                 ("108161779", "University  of   British Columbia"), ("108161779", "University-of-British-Columbia")]
        official = {"857305932": "Amazon.com.ca ULC", "108161779": "University of British Columbia"}
        for bn, name in cases:
            same_key = _name_key(name) == _name_key(official[bn])
            outcome = g.registry_lookup(bn, name, TODAY)
            self.assertEqual(outcome == "registered", same_key, f"{name!r}: CRA said {outcome}")


@unittest.skipUnless(LIVE, "set TAXCHECK_LIVE=1 to run live tests")
class LiveRegistrationDates(unittest.TestCase):
    """The day search (gst_dates.py) against the dates CRA publishes for simplified registrants."""

    @classmethod
    def setUpClass(cls):
        cls.gst = GstChecker(delay=1.0)
        cls.gst.simplified.lookup("000000000", TODAY)
        rows = cls.gst.simplified._rows or {}
        single = [(bn, es[0]) for bn, es in sorted(rows.items()) if len(es) == 1 and es[0]["registered"]]
        cls.active = next(((bn, e) for bn, e in single if not e["deregistered"]), None)
        cls.ended = next(((bn, e) for bn, e in single if e["deregistered"] and e["deregistered"] < TODAY), None)

    def test_start_date_matches_the_list(self):
        from taxcheck.gst_dates import registration_start

        # The registry answers per business number: a business that also had an ordinary GST/HST account
        # answers "registered" from that account's start, which can be years before its simplified one.
        bn, e = self.active
        start = registration_start(self.gst, bn, TODAY)
        self.assertIsNotNone(start)
        self.assertLessEqual(start, e["registered"])

    def test_cancellation_date_matches_the_list(self):
        from taxcheck.gst_dates import cancellation_date

        bn, e = self.ended
        registered_day = e["registered"] + (e["deregistered"] - e["registered"]) // 2
        self.assertEqual(cancellation_date(self.gst, bn, registered_day, TODAY), e["deregistered"])


@unittest.skipUnless(LIVE, "set TAXCHECK_LIVE=1 to run live tests")
class LiveDigitalEconomy(unittest.TestCase):
    def test_foreign_vendor_on_the_simplified_list(self):
        from taxcheck.checks import RowChecker

        rc = RowChecker(GstChecker(delay=1.0), None, available={"name", "gst", "country"})
        res = rc.check({"name": "Airalo", "gst": "91-0000001", "country": "US"}, TODAY)["gst"]
        self.assertEqual(res.status, REGISTERED, res.detail)
        self.assertEqual(res.facts["number_checked"], "751950577RT9999")


@unittest.skipUnless(LIVE, "set TAXCHECK_LIVE=1 to run live tests")
class LiveEndToEnd(unittest.TestCase):
    def test_sample_workbook_with_parallel_lookups_and_dates(self):
        import logging
        import shutil
        import tempfile
        from pathlib import Path

        import openpyxl

        import check_suppliers

        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "sample_suppliers.xlsx"
            shutil.copy(Path(__file__).parent.parent / "examples" / "sample_suppliers.xlsx", src)
            try:
                code = check_suppliers.main([str(src), "--workers", "3", "--gst-dates", "end", "--no-name-lookup"])
            finally:
                root = logging.getLogger()
                for h in list(root.handlers):
                    root.removeHandler(h)
                    h.close()
            self.assertEqual(code, 0)
            ws = openpyxl.load_workbook(Path(tmp) / "sample_suppliers_tax_check.xlsx").active
            headers = [c.value for c in ws[1]]
            rows = [dict(zip(headers, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
            statuses = [v for r in rows for k, v in r.items() if k and k.endswith(" Status")]
            self.assertNotIn("ERROR", statuses)
            amazon = next(r for r in rows if r["Supplier Name"] == "Amazon.com.ca ULC")
            self.assertEqual(amazon["GST/HST Registration End"], f"Active on {TODAY}")


if __name__ == "__main__":
    unittest.main()
