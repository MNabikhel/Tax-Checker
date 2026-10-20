"""Reading the supplier workbook and writing results back into a copy of it."""

import csv
import datetime as dt
import logging
import re
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .report import generic_action
from .result import (
    NAME_MATCHES, NAME_UNCONFIRMED, NAME_UPDATE,
    ERROR, INVALID, MANUAL, MISSING, NAME_NOT_MATCHED, NOT_APPLICABLE, NOT_CHECKED, NOT_CONFIRMED, NOT_LOOKED_UP, NOT_REGISTERED,
    REGISTERED,
)

log = logging.getLogger(__name__)

# Accepted header spellings (compared lowercase with spaces/punctuation removed).
COLUMNS = {
    "name": ["supplier name", "legal name", "legal entity", "legal entity name", "business name", "vendor name",
             "company name", "supplier", "vendor", "company", "name"],
    "trade_name": ["trade name", "operating name", "operating as", "dba", "dba name", "doing business as"],
    "province": ["province", "prov", "province/territory", "province / territory", "province/state", "state/prov",
                 "state/province", "region", "state"],
    "gst": ["gst/hst number", "gst/hst no", "gst/hst #", "gst/hst", "gst number", "gst no", "gst #", "gst", "hst number", "hst #", "hst",
            # ERP vendor export: "Tax ID" holds the GST/HST number for Canadian vendors
            # (and the US EIN or other foreign IDs for foreign vendors; see checks.RowChecker for routing).
            "tax id"],
    "bn": ["business number", "bn", "cra business number"],
    "qst": ["qst number", "qst no", "qst #", "qst", "tvq", "qst/tvq"],
    "bc_pst": ["bc pst number", "bc pst no", "bc pst #", "bc pst"],
    "sk_pst": ["sk pst number", "sk pst #", "sk pst", "saskatchewan pst"],
    "mb_rst": ["mb rst number", "mb rst #", "mb rst", "rst number", "rst #", "rst", "manitoba rst", "mb pst"],
    "pst": ["pst number", "pst no", "pst #", "pst", "provincial sales tax", "provincial sales tax number", "pst/rst"],
    "date": ["transaction date", "invoice date", "date",
             # ERP vendor export: date of the last cheque/payment item to the vendor.
             "check/ itm date", "check/itm date"],
    # Country and postal code decide whether Canadian checks apply and fill in a blank province.
    "country": ["country", "country code"],
    "postal": ["postal code", "postal", "postal/zip", "zip code", "zip"],
}

PROVINCES = {
    "AB": ["alberta", "alta", "alb"], "BC": ["british columbia", "bc"], "MB": ["manitoba", "man"],
    "NB": ["new brunswick"], "NL": ["newfoundland", "newfoundland and labrador", "nfld", "nf", "nfl", "lab"],
    "NS": ["nova scotia"], "NT": ["northwest territories", "nwt"], "NU": ["nunavut"],
    "ON": ["ontario", "ont"], "PE": ["prince edward island", "pei"], "QC": ["quebec", "québec", "que", "pq"],
    "SK": ["saskatchewan", "sask"], "YT": ["yukon", "yk"],
}

STATUS_FILL = {
    REGISTERED: "C6EFCE",
    NAME_NOT_MATCHED: "E2EFDA",
    NOT_REGISTERED: "FFC7CE",
    INVALID: "FFC7CE",
    NOT_CONFIRMED: "FFEB9C",
    MISSING: "FFEB9C",
    MANUAL: "DDEBF7",
    ERROR: "F8CBAD",
    NOT_APPLICABLE: "EDEDED",
    NOT_LOOKED_UP: "FFFFFF",
    NOT_CHECKED: "EDEDED",
}

NAME_CHECK_FILL = {"MATCHES": "C6EFCE", "UPDATE NAME": "FFEB9C", "NOT CONFIRMED": "F8CBAD"}

TAXES = [("gst", "GST/HST"), ("qst", "QST"), ("bc_pst", "BC PST"), ("sk_pst", "SK PST"), ("mb_rst", "MB RST")]


def _key(s):
    return re.sub(r"[^a-z0-9#]", "", str(s or "").lower())


_ALIASES, _ALIAS_RANK = {}, {}


def rebuild_aliases():
    """Header spelling -> field, and each spelling's position in its field's list: COLUMNS lists the specific
    spellings first ("postal code" before "zip"), so when two columns fit one field the more specific header wins.
    Rebuilt when a tailoring file adds spellings (taxcheck/tailoring.py)."""
    _ALIASES.clear()
    _ALIAS_RANK.clear()
    for col, names in COLUMNS.items():
        for i, a in enumerate(names):
            _ALIASES.setdefault(_key(a), col)
            _ALIAS_RANK.setdefault(_key(a), i)


rebuild_aliases()
TAX_FIELDS = {"gst", "bn", "qst", "pst", "bc_pst", "sk_pst", "mb_rst"}


def suggest_field(header):
    """For an unrecognized header, the closest known alias and its field, or None (used by --check-columns)."""
    import difflib

    key = _key(header)
    if not key:
        return None
    if re.search(r"\b(id|code|key|ref|no\.?)$", str(header).strip().lower()) and not re.search(
        r"tax|gst|hst|qst|tvq|pst|rst|bn|business", str(header).lower()
    ):
        return None  # "Vendor ID", "Supplier Code": identifiers, not names
    # Country/postal are matched by exact alias only: as fuzzy suggestions they misfire ("Amount" ~ "country").
    candidates = [a for a in _ALIASES if _ALIASES[a] not in ("country", "postal")]
    if re.search(r"tax|gst|hst|qst|tvq|pst|rst", key):
        # A tax-ish header should only be suggested as a tax-number field, never as province or name.
        candidates = [a for a in candidates if _ALIASES[a] in TAX_FIELDS]
    match = difflib.get_close_matches(key, candidates, n=1, cutoff=0.6)
    if not match:
        # Partial words: "Legal Entity" contains "legal"...
        match = [a for a in candidates if len(a) >= 3 and (a in key or key in a)][:1]
    return (match[0], _ALIASES[match[0]]) if match else None


COUNTRIES_CA = {"CA", "CAN", "CANADA", "CDN"}

# First letter of a Canadian postal code -> province (X is shared by NT and NU; NT is far more common).
POSTAL_PROVINCE = {
    "A": "NL", "B": "NS", "C": "PE", "E": "NB", "G": "QC", "H": "QC", "J": "QC", "K": "ON", "L": "ON",
    "M": "ON", "N": "ON", "P": "ON", "R": "MB", "S": "SK", "T": "AB", "V": "BC", "X": "NT", "Y": "YT",
}
CA_POSTAL = re.compile(r"^[ABCEGHJ-NPRSTVXY]\d[A-Z]\s?\d[A-Z]\d$")


def province_from_postal(value):
    """Province from a Canadian postal code ("T5J 0N3" -> AB), or "" if it isn't one."""
    s = str(value or "").strip().upper()
    return POSTAL_PROVINCE[s[0]] if CA_POSTAL.match(s) else ""


def normalize_country(value):
    """'CA', 'Canada' -> 'CA'; other values upper-cased as given; blank -> ''."""
    s = str(value or "").strip().upper().replace(".", "")
    return "CA" if s in COUNTRIES_CA else s


def normalize_province(value):
    """'BC', 'B.C.', 'Ont.', 'Newfoundland & Labrador' -> the two-letter code; anything else -> ""."""
    s = re.sub(r"\s+", " ", str(value or "").replace(".", "").replace("&", " and ")).strip()
    if s.upper() in PROVINCES:
        return s.upper()
    low = s.lower()
    for code, names in PROVINCES.items():
        if low in names:
            return code
    return ""


# Order matters for ambiguous dates: 09/10/2026 is read day-first (9 October), as in Canadian usage. A tailoring
# file can put other formats first (e.g. month-first "%m/%d/%Y" for an export that writes dates that way).
DATE_FORMATS = [
    "%Y-%m-%d", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y%m%d",
    "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y",
    "%d-%b-%Y", "%d %b %Y", "%d-%B-%Y", "%d %B %Y",
    "%B %d, %Y", "%b %d, %Y", "%B %d %Y", "%b %d %Y",
]


def parse_date(value):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    s = str(value or "").strip()
    for fmt in DATE_FORMATS:
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def as_text(value):
    """Text from outside sources is written as text: a leading '=' would make Excel treat it as a formula."""
    value = str(value or "")
    return "'" + value if value[:1] in ("=", "+", "@") else value


def header_hints(path, sheet=None):
    """When no header row is recognized: the most header-like row and a suggested field per cell.

    Used by --check-columns so it can still help when almost nothing matches.
    """
    path = Path(path)
    wb = _read_csv(path) if path.suffix.lower() == ".csv" else openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    best_row, best = None, []
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 15), values_only=True), 1):
        cells = [(c, v) for c, v in enumerate(row, 1) if isinstance(v, str) and v.strip()]
        if len(cells) > len(best):
            best_row, best = r, cells
    hints = []
    for c, header in best:
        field = _ALIASES.get(_key(header))
        hint = (header, field) if field else suggest_field(header)
        hints.append((get_column_letter(c), header, field, hint))
    return best_row, hints


def clean_value(value):
    """Excel stores typed numbers as floats (857305932 -> 857305932.0); turn those back into whole numbers."""
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


class InputError(ValueError):
    """A problem with the input file that the person running the script can fix."""


def _read_csv(path):
    """Load a CSV into a one-sheet workbook. Excel on Windows saves CSVs as cp1252, others as UTF-8."""
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            with open(path, newline="", encoding=encoding) as f:
                rows = list(csv.reader(f))
            break
        except UnicodeDecodeError:
            continue
    log.debug("Read CSV %s as %s: %d rows", path, encoding, len(rows))
    wb = openpyxl.Workbook()
    wb.active.title = re.sub(r"[\\/*?:\[\]]", "_", Path(path).stem)[:31] or "Sheet1"
    for row in rows:
        wb.active.append([v if v != "" else None for v in row])
    return wb


class SupplierSheet:
    def __init__(self, path, sheet=None):
        path = Path(path)
        if not path.exists():
            raise InputError(f"Can't find the input file: {path}")
        suffix = path.suffix.lower()
        if suffix == ".csv":
            self.wb = values_wb = _read_csv(path)
        elif suffix in (".xlsx", ".xlsm", ".xltx", ".xltm"):
            # Two copies: `wb` keeps formulas so the output is a faithful copy of the input;
            # `values_wb` has the calculated values (as last saved by Excel) that the checks read.
            self.wb = openpyxl.load_workbook(path)
            values_wb = openpyxl.load_workbook(path, data_only=True)
        elif suffix == ".xls":
            raise InputError(f"{path.name} is an old-style .xls file. Open it in Excel and use File > Save As > Excel Workbook (.xlsx).")
        else:
            raise InputError(f"{path.name}: expected an Excel workbook (.xlsx) or a .csv file.")
        name = sheet or self.wb.sheetnames[0]
        if name not in self.wb.sheetnames:
            raise InputError(f"There's no sheet called '{name}'. Sheets in this file: {', '.join(self.wb.sheetnames)}")
        self.ws, self.values_ws = self.wb[name], values_wb[name]
        self.header_row, self.cols = self._find_header()
        log.info("Reading sheet '%s': header on row %d", name, self.header_row)
        for col, c in self.cols.items():
            log.debug("  column %-10s <- '%s' (column %s)", col, self.ws.cell(self.header_row, c).value, get_column_letter(c))

    def _has_data(self, header_row, c, sample=200):
        last = min(self.values_ws.max_row, header_row + sample)
        return any(self.values_ws.cell(r, c).value not in (None, "") for r in range(header_row + 1, last + 1))

    def _find_header(self):
        for r in range(1, min(self.values_ws.max_row, 15) + 1):
            found = {}
            for c in range(1, self.values_ws.max_column + 1):
                key = _key(self.values_ws.cell(r, c).value)
                col = _ALIASES.get(key)
                if col:
                    found.setdefault(col, []).append((_ALIAS_RANK[key], c))
            if "name" in found and len(found) >= 2:
                cols = {}
                for col, options in found.items():
                    if len(options) == 1:
                        cols[col] = options[0][1]
                        continue
                    # Several columns fit (e.g. "Zip" and "Postal Code"): one with data beats an empty one, then
                    # the more specific spelling, then the leftmost.
                    best = min(options, key=lambda o: (not self._has_data(r, o[1]), o[0], o[1]))
                    cols[col] = best[1]
                    log.info("Columns %s all fit '%s'; using %s", ", ".join(
                        repr(self.values_ws.cell(r, c).value) for _, c in options), col,
                        repr(self.values_ws.cell(r, best[1]).value))
                return r, cols
        seen = [
            str(v).strip()
            for row in self.values_ws.iter_rows(min_row=1, max_row=min(self.values_ws.max_row, 5), values_only=True)
            for v in row
            if v not in (None, "")
        ]
        raise InputError(
            "Couldn't find a header row: it needs a supplier-name column plus at least one other recognized "
            f"column. Text in the first rows: {', '.join(repr(v) for v in seen[:20]) or '(none)'}. "
            "Run with --check-columns for suggestions, and add your header names to COLUMNS in taxcheck/workbook.py."
        )

    def describe_columns(self):
        """Lines explaining the column mapping, for --check-columns."""
        lines = [f"Header row: {self.header_row} on sheet '{self.ws.title}'", "Recognized columns:"]
        for col, c in self.cols.items():
            lines.append(f"  {get_column_letter(c):>3} '{self.values_ws.cell(self.header_row, c).value}' -> {col}")
        used = set(self.cols.values())
        others = []
        for c in range(1, self.values_ws.max_column + 1):
            header = self.values_ws.cell(self.header_row, c).value
            if c in used or header in (None, "") or header == "Overall":
                continue
            hint = suggest_field(header)
            known = _ALIASES.get(_key(header))
            if known:
                note = f"also matches '{known}', but column {get_column_letter(self.cols[known])} is used (a column with data, then the more specific header, then the leftmost)"
            elif hint:
                note = f"not recognized; closest alias '{hint[0]}' ({hint[1]}). If it's that field, add '{header}' to COLUMNS['{hint[1]}']"
            else:
                note = "not used"
            others.append(f"  {get_column_letter(c):>3} '{header}': {note}")
        if others:
            lines.append("Other columns:")
            lines.extend(others)
        missing = [col for col in ("gst", "province", "qst", "pst", "trade_name", "date") if col not in self.cols]
        if missing:
            lines.append(f"Fields with no column: {', '.join(missing)}")
        return lines

    def rows(self):
        """Yield (row_number, {column: value}) for every non-empty supplier row."""
        for r in range(self.header_row + 1, self.values_ws.max_row + 1):
            values = {col: clean_value(self.values_ws.cell(r, c).value) for col, c in self.cols.items()}
            if any(v not in (None, "") for v in values.values()):
                yield r, values

    def write_results(self, results, out_path, run_notes, ok_statuses=(REGISTERED,)):
        """results: {row_number: {tax_key: Result}}. Adds the result columns to the supplier sheet (a copy: the
        input file is never saved over), an "Action List" sheet and a "Tax Check Summary" sheet."""
        from .report import (SEV_LATE, SEV_NAME, SEV_NAME_UNCONFIRMED, SEV_OK, SEV_OTHER_NAME, SEVERITY_FILL,
                             issue_for, name_source, number_registered, row_summary)

        ws = self.ws
        # Re-running on an earlier results file overwrites its result columns instead of adding more.
        start = next(
            (c for c in range(1, ws.max_column + 1) if ws.cell(self.header_row, c).value == "Overall"),
            None,
        )
        rerun = start is not None
        start = start or ws.max_column + 1
        # Only the taxes chosen for this run get columns (self.taxes is set by check_suppliers; None = all).
        chosen = getattr(self, "taxes", None)
        taxes = [(k, label) for k, label in TAXES if chosen is None or k == "gst" or k in chosen]
        labels = dict(TAXES)
        for name in ("Tax Check Summary", "Action List"):
            if name in self.wb.sheetnames:
                del self.wb[name]
        if rerun:  # an earlier run may have had more tax columns: clear all of its result columns first
            for row in ws.iter_rows(min_row=self.header_row, min_col=start, max_col=ws.max_column):
                for cell in row:
                    cell.value, cell.fill = None, PatternFill()

        # (header, width, kind) per column; kind decides how the value is found for a row.
        columns = [
            ("Overall", 16, "overall"), ("Main Issue", 40, "main"), ("Action Needed", 60, "action"),
            ("Other Issues", 40, "other"), ("Vendor Name Check", 18, "name_check"),
            ("Suggested Vendor Name", 34, "suggested"), ("Do-Not-Use Vendor", 12, "dnu"),
            ("Province Used", 14, "province"),
        ]
        for key, label in taxes:
            if key == "gst":
                columns += [
                    (f"{label} Status", 18, ("status", key)),
                    (f"{label} Number Registered", 22, ("number_registered", key)),
                    (f"{label} Issue", 36, ("issue", key)),
                    (f"{label} Tax ID Problem", 30, ("fact", key, "tax_id_problem")),
                    (f"{label} Number Checked", 18, ("fact", key, "number_checked")),
                    (f"{label} Date Checked", 13, ("fact", key, "date_checked")),
                    (f"{label} Registered Since", 14, ("fact", key, "registered_since")),
                    (f"{label} Registration End", 18, ("fact", key, "registration_end")),
                    (f"{label} Status Today", 18, ("fact", key, "status_today")),
                    (f"{label} Name Accepted by CRA", 34, ("fact", key, "accepted_name")),
                    (f"{label} Name Source", 26, ("source", key)),
                    (f"{label} Registry Name", 34, ("fact", key, "official_name")),
                    (f"{label} Details", 55, ("detail", key)),
                ]
            else:
                columns += [
                    (f"{label} Status", 18, ("status", key)), (f"{label} Issue", 32, ("issue", key)),
                    (f"{label} Number Checked", 18, ("fact", key, "number_checked")),
                    (f"{label} Registered Since", 14, ("fact", key, "registered_since")),
                    (f"{label} Registration End", 18, ("fact", key, "registration_end")),
                    (f"{label} Name on Record", 30, ("registered_name", key)),
                    (f"{label} Details", 50, ("detail", key)),
                ]
        columns.append(("Checked On", 12, "checked_on"))

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="104735")
        for i, (h, width, _) in enumerate(columns):
            cell = ws.cell(self.header_row, start + i, h)
            cell.font, cell.fill = header_font, header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.column_dimensions[get_column_letter(start + i)].width = width
        today = dt.date.today().isoformat()

        counts = {key: Counter() for key, _ in taxes}
        issue_counts = {key: Counter() for key, _ in taxes}
        issue_actions = {}
        overall_counts, name_counts = Counter(), Counter()
        action_rows = []
        wrap = Alignment(wrap_text=True, vertical="top")
        shared, cut_off, name_limit = self._cross_row_notes(results)
        for r, by_tax in results.items():
            # (An earlier run's result columns were already cleared above, once, when re-running on its output.)
            statuses = [by_tax[k].status for k, _ in taxes if k in by_tax]
            if statuses and all(s == NOT_APPLICABLE for s in statuses):
                overall = "NOT APPLICABLE"  # foreign supplier, nothing Canadian to check
            elif statuses and all(s == NOT_CHECKED for s in statuses):
                overall = "NOT CHECKED"  # --paid-since: last paid before the cut-off
            elif statuses and all(s == NOT_LOOKED_UP for s in statuses):
                overall = "PENDING"  # --offline: nothing wrong found offline, still needs the online run
            elif statuses and all(s in ok_statuses or s == NOT_APPLICABLE for s in statuses):
                overall = "OK"
            else:
                overall = "REVIEW"
            summary_row = row_summary(by_tax, taxes)
            # e.g. registered on the payment date but cancelled since. With --trust-gst-number (NAME NOT MATCHED in
            # ok_statuses) a name that differs from CRA's is accepted, so only other problems count here.
            worst = max((issue_for(k, by_tax[k])[2] for k, _ in taxes if k in by_tax
                         and not (by_tax[k].status == NAME_NOT_MATCHED and NAME_NOT_MATCHED in ok_statuses
                                  and issue_for(k, by_tax[k])[2] in (SEV_OTHER_NAME, SEV_NAME_UNCONFIRMED))),
                        default=SEV_OK)
            if overall == "OK" and worst >= SEV_LATE:
                overall = "REVIEW"
            overall_counts[overall] += 1
            gst = by_tax.get("gst")
            gf = gst.facts if gst is not None else {}
            name_check = getattr(gst, "name_check", "") if gst is not None else ""
            if name_check:
                name_counts[name_check] += 1
            suggested = ""
            if gst is not None and (name_check == NAME_UPDATE or (gst.status == NAME_NOT_MATCHED and gf.get("accepted_name"))):
                suggested = gf.get("accepted_name") or gst.registered_name
            issues = {}
            for key, _ in taxes:
                if key in by_tax:
                    issues[key] = issue_for(key, by_tax[key])
                    counts[key][by_tax[key].status] += 1
                    issue_counts[key][issues[key][0]] += 1
                    issue_actions.setdefault((key, issues[key][0]), generic_action(issues[key][0], issues[key][1]))

            if r in shared:
                issue_counts["gst"][self.SHARED_NUMBER] += 1
                issue_actions.setdefault(("gst", self.SHARED_NUMBER), self.SHARED_ACTION)
                summary_row["other"] = "; ".join(filter(None, [summary_row["other"], *shared[r]]))
                summary_row["severity"] = max(summary_row["severity"], SEV_NAME)
            if r in cut_off:
                summary_row["action"] = " ".join(filter(None, [summary_row["action"], (
                    f"The vendor name is {name_limit} characters, the longest the vendor master holds here, so it's "
                    "probably cut off: CRA only accepts the full legal name.")]))
            for i, (_, _, kind) in enumerate(columns):
                cell = ws.cell(r, start + i)
                value, fill = None, None
                if kind == "overall":
                    value = overall
                    fill = {"OK": "C6EFCE", "NOT APPLICABLE": "EDEDED", "PENDING": "FFFFFF", "NOT CHECKED": "EDEDED"}.get(overall, "FFEB9C")
                elif kind == "main":
                    value, fill = summary_row["main"], SEVERITY_FILL.get(summary_row["severity"])
                elif kind == "action":
                    value = summary_row["action"]
                elif kind == "other":
                    value = summary_row["other"]
                elif kind == "name_check":
                    value, fill = name_check, NAME_CHECK_FILL.get(name_check)
                elif kind == "suggested":
                    value = suggested
                elif kind == "dnu":
                    value = "Yes" if gf.get("dnu") else None
                elif kind == "province":
                    value = gf.get("province_used") or None
                elif kind == "checked_on":
                    value = today
                else:
                    what, key = kind[0], kind[1]
                    res = by_tax.get(key)
                    if res is None:
                        continue
                    if what == "status":
                        value, fill = res.status, STATUS_FILL.get(res.status)
                    elif what == "number_registered":
                        value = number_registered(res) or None
                    elif what == "issue":
                        value, fill = issues[key][0], SEVERITY_FILL.get(issues[key][2])
                    elif what == "fact":
                        value = res.facts.get(kind[2]) or None
                    elif what == "source":
                        value = name_source(res) or None
                    elif what == "registered_name":
                        value = res.registered_name or None
                    elif what == "detail":
                        value = res.detail
                if value in (None, ""):
                    continue
                cell.value = as_text(value) if isinstance(value, str) else value
                if fill:
                    cell.fill = PatternFill("solid", fgColor=fill)
                if kind in ("action", "other") or (isinstance(kind, tuple) and kind[0] == "detail"):
                    cell.alignment = wrap
            if overall == "REVIEW" or summary_row["severity"] >= SEV_NAME:
                action_rows.append((summary_row["severity"], r, overall, summary_row, suggested, by_tax, gf))

        name_col = self.cols.get("name", 1)
        ws.freeze_panes = ws.cell(self.header_row + 1, name_col + 1)  # keep the supplier name in view
        # Filter buttons on the header row, so each situation can be picked out in a large vendor list.
        ws.auto_filter.ref = f"A{self.header_row}:{get_column_letter(start + len(columns) - 1)}{ws.max_row}"

        # Sheet order: suppliers, Action List, Tax Check Summary.
        self._write_summary(today, taxes, counts, issue_counts, issue_actions, overall_counts, name_counts,
                            run_notes, header_font, header_fill)
        self._write_action_list(action_rows, taxes, header_font, header_fill, wrap)
        self.wb.save(out_path)

    SHARED_NUMBER = "Same GST/HST number as another vendor"
    SHARED_ACTION = ("Check whether these vendor records are the same business (a duplicate, or a trade name) or whether "
                     "one of them has the wrong GST/HST number; see 'Other Issues' for the other vendor.")

    def _cross_row_notes(self, results):
        """Findings that need the whole sheet. Returns ({row: [notes]} for GST/HST numbers that differently named
        vendors share, {rows whose name is probably cut off}, the vendor master's apparent name length limit)."""
        from .checks import clean_vendor_name, core_name

        names = {r: str(self._cell_value(r, "name") or "").strip() for r in results}
        by_number = defaultdict(list)
        for r, by_tax in results.items():
            gst = by_tax.get("gst")
            # Only numbers CRA knows: an invalid or placeholder number shared by many rows isn't a finding.
            if gst is None or gst.status not in (REGISTERED, NAME_NOT_MATCHED, NOT_REGISTERED):
                continue
            number = re.sub(r"\D", "", str(gst.facts.get("number_checked") or ""))[:9]
            if len(number) == 9 and names[r]:
                by_number[number].append(r)
        shared = defaultdict(list)
        for rows in by_number.values():
            # The same vendor spelt two ways ("ACME" / "ACME INC.", "Café Ltée" / "Cafe Ltee") isn't flagged.
            keys = {r: core_name(clean_vendor_name(names[r])[0]) or names[r].lower() for r in rows}
            if len(set(keys.values())) < 2:
                continue
            first_of = {}  # one example row per distinct name, so long groups stay quick
            for r in rows:
                first_of.setdefault(keys[r], r)
            for r in rows:
                others = [o for k, o in first_of.items() if k != keys[r]]
                text = []
                for o in others[:3]:
                    gst = results[o]["gst"]
                    accepted = gst.facts.get("accepted_name") if gst.status == REGISTERED else ""
                    text.append(f"row {o} '{names[o]}'" + (f" (CRA accepts the number under '{accepted}')" if accepted else ""))
                more = f" and {len(others) - 3} more" if len(others) > 3 else ""
                shared[r].append("GST/HST: same number as " + "; ".join(text) + more)

        # A vendor-master name field has a length limit (40 characters in one ERP): names that long are probably
        # cut off, and CRA only accepts the full legal name. Detected from the sheet itself: the longest length,
        # if several names reach it exactly.
        lengths = Counter(len(n) for n in names.values() if n)
        limit = max(lengths) if lengths else 0
        cut = set()
        if limit >= 30 and lengths[limit] >= 3:
            cut = {r for r, n in names.items() if len(n) == limit and results[r].get("gst") is not None
                   and getattr(results[r]["gst"], "name_check", "") == NAME_UNCONFIRMED}
        return shared, cut, limit

    def _cell_value(self, r, field):
        c = self.cols.get(field)
        return self.values_ws.cell(r, c).value if c else None

    def _write_action_list(self, action_rows, taxes, header_font, header_fill, wrap):
        """One row per supplier that needs something done, worst first: the work list."""
        from .report import SEVERITY_FILL

        sheet = self.wb.create_sheet("Action List", 1)  # created after the summary, so it lands before it
        first_header = self.values_ws.cell(self.header_row, 1).value
        id_col = 1 if self.cols.get("name") != 1 else None
        headers = (([str(first_header or "ID")] if id_col else []) +
                   ["Company Name", "Province Used", "Tax ID", "Tax ID Problem", "Overall", "Main Issue", "Action Needed",
                    "Suggested Vendor Name", "Other Issues"] + [f"{label} Status" for _, label in taxes] + ["Sheet Row"])
        sheet.append(headers)
        widths = ([16] if id_col else []) + [38, 14, 18, 30, 12, 40, 70, 34, 40] + [18] * len(taxes) + [10]
        for i, (cell, width) in enumerate(zip(sheet[1], widths), 1):
            cell.font, cell.fill, cell.alignment = header_font, header_fill, Alignment(wrap_text=True, vertical="top")
            sheet.column_dimensions[get_column_letter(i)].width = width
        action_rows.sort(key=lambda x: (-x[0], x[3]["main"], x[1]))
        for sev, r, overall, summary_row, suggested, by_tax, gf in action_rows:
            values = ([self.values_ws.cell(r, id_col).value] if id_col else []) + [
                self._cell_value(r, "name"), gf.get("province_used"), self._cell_value(r, "gst"),
                gf.get("tax_id_problem") or None, overall,
                summary_row["main"], summary_row["action"], suggested, summary_row["other"],
            ] + [by_tax[k].status if k in by_tax else None for k, _ in taxes] + [r]
            sheet.append([as_text(v) if isinstance(v, str) else v for v in values])
            row = sheet.max_row
            main_col = headers.index("Main Issue") + 1
            sheet.cell(row, main_col).fill = PatternFill("solid", fgColor=SEVERITY_FILL.get(sev, "FFFFFF"))
            sheet.cell(row, headers.index("Action Needed") + 1).alignment = wrap
        sheet.freeze_panes = "B2"
        if sheet.max_row > 1:
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sheet.max_row}"

    def _write_summary(self, today, taxes, counts, issue_counts, issue_actions, overall_counts, name_counts,
                       run_notes, header_font, header_fill):
        summary = self.wb.create_sheet("Tax Check Summary", 1)
        bold = Font(bold=True)

        def heading(values):
            summary.append(values)
            for cell in summary[summary.max_row]:
                if cell.value is not None:
                    cell.font, cell.fill = header_font, header_fill

        summary.append(["Supplier tax registration check", today])
        summary["A1"].font = Font(bold=True, size=14)
        summary.append([])
        heading(["Overall", "OK", "REVIEW", "NOT APPLICABLE", "PENDING", "NOT CHECKED", "Rows checked"])
        summary.append(["Suppliers"] + [overall_counts[k] for k in ("OK", "REVIEW", "NOT APPLICABLE", "PENDING", "NOT CHECKED")]
                       + [sum(overall_counts.values())])
        summary.append([])
        heading(["Vendor Name Check", NAME_MATCHES, NAME_UPDATE, NAME_UNCONFIRMED])
        summary.append(["Suppliers"] + [name_counts[k] for k in (NAME_MATCHES, NAME_UPDATE, NAME_UNCONFIRMED)])
        summary.append(["MATCHES: the vendor name is the name CRA has. UPDATE NAME: registration confirmed, but under "
                        "the name in 'Suggested Vendor Name'; correct the vendor record. NOT CONFIRMED: CRA accepted no "
                        "name we have; get the legal name from the supplier (e.g. their invoice)."])
        summary.append(["Registration dates: QST dates come from Revenu Québec; GST/HST dates from CRA's simplified "
                        "list, or by asking the CRA registry about specific days (the registry gives no dates "
                        "itself). 'Registration End' is when a registration was cancelled; 'Active on <date>' means "
                        "still registered that day. Registrations don't expire. Manitoba, BC and Saskatchewan's "
                        "registries don't give dates."])
        summary.append([])
        heading(["Issues by tax", "Issue", "Suppliers", "What to do"])
        for key, label in taxes:
            for issue, n in sorted(issue_counts[key].items(), key=lambda x: -x[1]):
                summary.append([label, issue, n, issue_actions.get((key, issue), "")])
        summary.append([])
        statuses = [REGISTERED, NAME_NOT_MATCHED, NOT_REGISTERED, NOT_CONFIRMED, INVALID, MISSING, MANUAL, ERROR,
                    NOT_APPLICABLE, NOT_LOOKED_UP, NOT_CHECKED]
        heading(["Status by tax"] + statuses)
        for key, label in taxes:
            summary.append([label] + [counts[key][s] for s in statuses])
        summary.append([])
        heading(["Notes"])
        for note in run_notes:
            summary.append([note])
        summary.column_dimensions["A"].width = 30
        summary.column_dimensions["B"].width = 42
        summary.column_dimensions["C"].width = 12
        summary.column_dimensions["D"].width = 90
